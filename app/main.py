from __future__ import annotations

import hashlib
import io
import json
import mimetypes
import os
import re
import shutil
import sqlite3
import threading
import uuid
import zipfile
from copy import deepcopy
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from PIL import Image, ImageFilter
from pydantic import BaseModel, Field, ValidationError

from .file_handlers import (
    SUPPORTED,
    content_manifest,
    extension,
    extract_text,
    mask_file,
)
from .ner import MODEL_ID, RaanerService
from .ocr import OCRService
from .policies import DEFAULT_POLICIES
from .security import decrypt_json, encrypt_json

ROOT=Path(__file__).resolve().parent.parent; DATA=ROOT/'data'; TASKS=DATA/'tasks'; MODELS=ROOT/'models'
for p in (DATA,TASKS,MODELS): p.mkdir(parents=True, exist_ok=True)
DB=DATA/'tasks.db'; MAX_UPLOAD=50*1024*1024; MAX_BATCH_FILES=int(os.getenv('MAX_BATCH_FILES','20')); MAX_BATCH_BYTES=int(os.getenv('MAX_BATCH_BYTES',str(500*1024*1024))); TASK_TTL_HOURS=int(os.getenv('TASK_TTL_HOURS','24')); jobs={}; batches={}; policy_store=dict(DEFAULT_POLICIES); ner_service=RaanerService(MODELS/'raner'); ocr_service=OCRService(MODELS/'paddleocr')
TASK_ID_RE = re.compile(r'^[0-9a-f]{32}$', re.IGNORECASE)
IMAGE_EXTENSIONS = {'png', 'jpg', 'jpeg', 'bmp', 'tif', 'tiff'}
PUBLIC_ARTIFACT_NAMES = {'analysis.json', 'report.json', 'result.md', 'masked.md', 'content.json', 'entities.json', 'masked.png', 'original.png', 'results.zip'}
app=FastAPI(title='本地数据脱敏系统',version='0.2.0')

def conn():
    c=sqlite3.connect(DB,check_same_thread=False, timeout=10)
    c.execute('PRAGMA journal_mode=WAL'); c.execute('PRAGMA busy_timeout=10000'); c.execute('PRAGMA foreign_keys=ON')
    c.execute('CREATE TABLE IF NOT EXISTS tasks(id TEXT PRIMARY KEY, original TEXT, masked TEXT, created REAL, kind TEXT DEFAULT "text")')
    if 'kind' not in {row[1] for row in c.execute('PRAGMA table_info(tasks)').fetchall()}:
        c.execute('ALTER TABLE tasks ADD COLUMN kind TEXT DEFAULT "text"')
    columns = {row[1] for row in c.execute('PRAGMA table_info(tasks)').fetchall()}
    migrations = {
        'original_name': 'TEXT DEFAULT ""', 'source_path': 'TEXT DEFAULT ""',
        'artifact_name': 'TEXT DEFAULT ""', 'source_sha256': 'TEXT DEFAULT ""',
        'status': 'TEXT DEFAULT "created"', 'reversible': 'INTEGER DEFAULT 0',
        'batch_id': 'TEXT DEFAULT ""', 'metadata_json': 'TEXT DEFAULT "{}"',
        'updated': 'REAL DEFAULT 0', 'error': 'TEXT DEFAULT ""',
    }
    for name, definition in migrations.items():
        if name not in columns:
            c.execute(f'ALTER TABLE tasks ADD COLUMN {name} {definition}')
    c.execute('CREATE TABLE IF NOT EXISTS rules(id TEXT PRIMARY KEY,name TEXT,kind TEXT,pattern TEXT,entity_type TEXT,replacement TEXT,enabled INTEGER DEFAULT 1)')
    c.execute('''CREATE TABLE IF NOT EXISTS batches(
        id TEXT PRIMARY KEY, created REAL NOT NULL, updated REAL NOT NULL,
        status TEXT NOT NULL, total INTEGER DEFAULT 0, completed INTEGER DEFAULT 0,
        failed INTEGER DEFAULT 0, archive_name TEXT DEFAULT '', metadata_json TEXT DEFAULT '{}',
        error TEXT DEFAULT '')''')
    c.commit(); return c

def _valid_task_id(value: str) -> bool:
    return bool(isinstance(value, str) and TASK_ID_RE.fullmatch(value))

def _task_path(task_id: str, *, create: bool = False) -> Path:
    if not _valid_task_id(task_id):
        raise HTTPException(404, 'task not found')
    base = TASKS.resolve(); path = (base / task_id).resolve()
    if path.parent != base:
        raise HTTPException(404, 'task not found')
    if create:
        path.mkdir(parents=True, exist_ok=True)
    return path

def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()

def _safe_filename(filename: str) -> str:
    name = Path(filename or 'upload.bin').name.replace('\x00', '')
    if not name or name in {'.', '..'}:
        return 'upload.bin'
    return name[:240]

def _validate_upload(data: bytes, filename: str, ext: str) -> None:
    """Validate the file signature before handing bytes to a parser."""
    if ext in {'txt', 'md'}:
        if b'\x00' in data[:4096]:
            raise HTTPException(415, '文本文件包含二进制内容')
        return
    if ext == 'pdf':
        if not data.lstrip().startswith(b'%PDF-'):
            raise HTTPException(415, 'PDF 文件签名无效')
        try:
            import fitz
            document = fitz.open(stream=data, filetype='pdf')
            if document.page_count > 1000:
                raise HTTPException(413, 'PDF 页数超过 1000 页')
            document.close()
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(415, f'无法读取 PDF: {exc}') from exc
        return
    if ext in {'docx', 'xlsx', 'xlsm'}:
        if not data.startswith(b'PK\x03\x04'):
            raise HTTPException(415, 'Office 文件签名无效')
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                infos = archive.infolist()
                if len(infos) > 10000 or sum(max(0, item.file_size) for item in infos) > 500 * 1024 * 1024:
                    raise HTTPException(413, 'Office 压缩包内容过大')
                if any(item.flag_bits & 0x1 for item in infos):
                    raise HTTPException(415, '不支持加密的 Office 文件')
                names = set(archive.namelist())
                if ext == 'docx' and 'word/document.xml' not in names:
                    raise HTTPException(415, 'DOCX 结构无效')
                if ext in {'xlsx', 'xlsm'} and 'xl/workbook.xml' not in names:
                    raise HTTPException(415, 'Excel 结构无效')
        except HTTPException:
            raise
        except (OSError, zipfile.BadZipFile) as exc:
            raise HTTPException(415, f'Office 文件损坏: {exc}') from exc
        return
    if ext in IMAGE_EXTENSIONS:
        try:
            with Image.open(io.BytesIO(data)) as image:
                image.verify()
        except Exception as exc:
            raise HTTPException(415, f'图片文件损坏: {exc}') from exc
        return
    raise HTTPException(415, f'不支持的文件格式: .{ext}')
def cleanup():
    c=conn(); cutoff=(datetime.now()-timedelta(hours=TASK_TTL_HOURS)).timestamp(); rows=c.execute('SELECT id FROM tasks WHERE created<?',(cutoff,)).fetchall()
    for (i,) in rows:
        if _valid_task_id(i):
            shutil.rmtree(TASKS / i, ignore_errors=True)
    c.execute('DELETE FROM tasks WHERE created<?',(cutoff,)); c.execute('DELETE FROM batches WHERE created<?',(cutoff,)); c.commit(); c.close()
def valid_id(v):
    if not re.fullmatch(r'\d{17}[\dXx]',v): return False
    try: datetime.strptime(v[6:14],'%Y%m%d')
    except ValueError: return False
    return '10X98765432'[sum(int(x)*w for x,w in zip(v[:17],[7,9,10,5,8,4,2,1,6,3,7,9,10,5,8,4,2]))%11]==v[-1].upper()
def luhn(v):
    return sum((n if i%2==0 else (n*2-9 if n>4 else n*2)) for i,n in enumerate(map(int,v[::-1])))%10==0
def valid_ip(v):
    parts = v.split('.')
    return len(parts) == 4 and all(re.fullmatch(r'[0-9]{1,3}', part) and 0 <= int(part) <= 255 for part in parts)
PATTERNS=[('PHONE',r'(?<!\d)(?:1[3-9]\d{9}|0\d{2,3}-?\d{7,8})(?!\d)',None),('ID_CARD',r'(?<!\w)\d{17}[\dXx](?!\w)',valid_id),('BANK_CARD',r'(?<!\d)\d{16,19}(?!\d)',luhn),('EMAIL',r'[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}',None),('IP_ADDRESS',r'(?<!\d)(?:\d{1,3}\.){3}\d{1,3}(?!\d)',valid_ip),('DATE_TIME',r'\d{4}[年\-/]\d{1,2}[月\-/]\d{1,2}日?(?:\s+\d{1,2}:\d{2}(?::\d{2})?)?',None),('MONEY',r'(?:人民币|RMB|￥|¥)\s?\d+(?:\.\d+)?|\d+(?:\.\d+)?(?:元|万元|亿元|美元|欧元)',None),('LICENSE_PLATE',r'[京津沪渝冀豫云辽黑湘皖鲁新苏浙赣鄂桂甘晋蒙陕吉闽贵粤青藏川宁琼][A-Z][A-Z0-9挂学警港澳领使]{5,6}',None)]
def analyze(text):
    c=conn(); custom=c.execute('SELECT name,kind,pattern,entity_type FROM rules WHERE enabled=1').fetchall(); c.close(); found=[]
    for typ,pat,check in PATTERNS+[(x[3],x[2],None) for x in custom if x[1]=='regex']:
        for m in re.finditer(pat,text,re.IGNORECASE):
            if check and not check(m.group()): continue
            found.append({'id':uuid.uuid4().hex,'type':typ,'text':m.group(),'start':m.start(),'end':m.end(),'score':.99 if check else .95,'recognizer':'builtin','source':'regex','selected':True,'box_ids':[]})
    for name,kind,pattern,typ in custom:
        if kind=='word':
            for m in re.finditer(re.escape(pattern),text,re.IGNORECASE): found.append({'id':uuid.uuid4().hex,'type':typ,'text':m.group(),'start':m.start(),'end':m.end(),'score':1.0,'recognizer':'custom','source':'custom','selected':True,'box_ids':[]})
    found.extend([{**e,'id':uuid.uuid4().hex,'selected':True,'box_ids':[]} for e in ner_service.analyze(text)])
    found.sort(key=lambda x:(x['start'],-x['end'])); accepted=[]
    for e in found:
        if not any(e['start']<a['end'] and a['start']<e['end'] for a in accepted): accepted.append(e)
    # Normalize every detection through Presidio's public result type so all
    # adapters expose one protocol regardless of recognizer implementation.
    try:
        from presidio_analyzer import RecognizerResult
        normalized = []
        for entity in accepted:
            result = RecognizerResult(entity_type=entity['type'], start=entity['start'], end=entity['end'], score=entity.get('score', 0.0))
            normalized.append({**entity, 'type': result.entity_type, 'start': result.start, 'end': result.end, 'score': result.score})
        return normalized
    except ImportError:
        return accepted
def register_task(task_id, kind, original='', masked='', *, source_path='', artifact_name='',
                  original_name=None, source_sha256='', status='created', reversible=False,
                  batch_id='', metadata=None):
    """Register task metadata without putting newly uploaded content in SQLite."""
    if not _valid_task_id(task_id):
        raise ValueError('invalid task id')
    now = datetime.now().timestamp()
    # ``original``/``masked`` are retained only for reading pre-migration legacy
    # rows. New callers pass the file paths and leave both content columns empty.
    if original_name is None:
        original_name = original if kind != 'text' else ''
    c=conn()
    c.execute('''INSERT OR REPLACE INTO tasks
        (id,original,masked,created,kind,original_name,source_path,artifact_name,
         source_sha256,status,reversible,batch_id,metadata_json,updated,error)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''',
        (task_id, original if kind in {'legacy_text','text'} else '',
         masked if kind in {'legacy_text','text'} else '', now, kind, original_name or '',
         source_path or '', artifact_name or '', source_sha256 or '', status,
         int(bool(reversible)), batch_id or '', json.dumps(metadata or {}, ensure_ascii=False), now, ''))
    c.commit(); c.close(); return task_id

def update_task(task_id, masked='', *, status='completed', artifact_name=None,
                source_path=None, reversible=None, error=''):
    if not _valid_task_id(task_id):
        raise ValueError('invalid task id')
    c=conn(); fields=['masked=?','status=?','updated=?','error=?']; values=[masked or '',status,datetime.now().timestamp(),error or '']
    if artifact_name is not None: fields.append('artifact_name=?'); values.append(artifact_name)
    if source_path is not None: fields.append('source_path=?'); values.append(source_path)
    if reversible is not None: fields.append('reversible=?'); values.append(int(bool(reversible)))
    values.append(task_id); c.execute(f"UPDATE tasks SET {', '.join(fields)} WHERE id=?", values); c.commit(); c.close()

def store(original,masked,kind='text'):
    i=uuid.uuid4().hex; d=_task_path(i, create=True); (d/'original.txt').write_text(original,encoding='utf-8'); (d/'masked.txt').write_text(masked,encoding='utf-8')
    register_task(i, kind, source_path='original.txt', artifact_name='masked.txt',
                  original_name='text.txt', source_sha256=_sha256(original.encode('utf-8')),
                  metadata={'legacy_text': True})
    return i

class TextIn(BaseModel): text:str=Field(min_length=1,max_length=2_000_000)
class MaskIn(BaseModel): text:str; entities:list[dict[str,Any]]
class RestoreIn(BaseModel): task_id:str; masked_text:str|None=None
class RestoreFileIn(BaseModel):
    password: str = Field(min_length=1, max_length=256)
class RuleIn(BaseModel): name:str; kind:str; pattern:str; entity_type:str='CUSTOM'; replacement:str='__MASKED_CUSTOM__'
class Box(BaseModel): id:str=Field(default_factory=lambda:uuid.uuid4().hex); page:int=1; x:float; y:float; width:float; height:float; text:str=''; selected:bool=True; source:str='manual'
class ImageMaskIn(BaseModel): analysis_id:str; boxes:list[Box]=[]; blur_radius:int=Field(default=14,ge=1,le=80)
class DocumentMaskIn(MaskIn): analysis_id:str|None=None

def apply_mask(text, entities):
    masked=text; seed=uuid.uuid4().hex
    for e in sorted((e for e in entities if e.get('selected',True)),key=lambda x:x['start'],reverse=True):
        if text[e['start']:e['end']]!=e['text']: raise HTTPException(422,'实体区间与原文不一致')
        token=f"__MASKED_{e['type']}_{hashlib.sha1((seed+e['text']+e['type']).encode()).hexdigest()[:8].upper()}__"
        masked=masked[:e['start']]+token+masked[e['end']:]
    return masked

@app.middleware('http')
async def housekeeping(request,call_next): cleanup(); return await call_next(request)
@app.get('/api/health')
def health(): return {'status':'ok','local_only':True,'version':app.version}
@app.get('/api/models/status')
def model_status():
    return {'initialized':ner_service.available, 'ner_available':ner_service.available,
            'ocr_available':ocr_service.engine is not None, 'device':'cpu', 'model':MODEL_ID,
            'ner_error':ner_service.error, 'ocr_error':ocr_service.error,
            'status': 'ready' if ner_service.available and ocr_service.engine else ('partial' if ner_service.available or ocr_service.engine else 'unavailable')}
@app.post('/api/models/initialize')
def initialize():
    i=uuid.uuid4().hex; jobs[i]={'status':'running','progress':5,'message':'正在初始化中文模型'}
    def run():
        try:
            jobs[i].update({'progress':20, 'ner_status':'running'})
            ner_service.load(download=True, device='cpu')
            jobs[i].update({'progress':60, 'ner_status':'ready' if ner_service.available else 'error', 'ner_error':ner_service.error})
            ocr_service.load()
            ready = ner_service.available or ocr_service.engine
            jobs[i]={'status':'completed' if ready else 'error','progress':100 if ready else 0,
                     'message':'模型初始化完成' if ready else '模型初始化失败',
                     'ner_status':'ready' if ner_service.available else 'error',
                     'ocr_status':'ready' if ocr_service.engine else 'error',
                     'ner_error':ner_service.error, 'ocr_error':ocr_service.error}
        except Exception as exc:
            jobs[i]={'status':'error','progress':0,'message':f'模型初始化异常: {exc}',
                     'ner_status':'ready' if ner_service.available else 'error',
                     'ocr_status':'ready' if ocr_service.engine else 'error', 'ner_error':ner_service.error, 'ocr_error':str(exc)}
    threading.Thread(target=run,daemon=True).start(); return {'job_id':i}
@app.get('/api/models/jobs/{job_id}')
def model_job(job_id): return jobs.get(job_id,{'status':'not_found'})
@app.post('/api/text/analyze')
def text_analyze(v:TextIn): return {'analysis_id':uuid.uuid4().hex,'text':v.text,'entities':analyze(v.text)}
@app.post('/api/text/mask')
def text_mask(v:MaskIn):
    masked=apply_mask(v.text,v.entities)
    return {'task_id':store(v.text,masked),'masked_text':masked}
@app.post('/api/text/unmask')
def text_unmask(v:RestoreIn):
    if not _valid_task_id(v.task_id): raise HTTPException(404,'task not found')
    c=conn(); row=c.execute('SELECT original,masked,source_path,artifact_name FROM tasks WHERE id=?',(v.task_id,)).fetchone(); c.close()
    if not row: raise HTTPException(404,'task not found')
    original, masked, source_path, artifact_name = row
    folder = _task_path(v.task_id)
    if source_path and (folder / source_path).is_file():
        original = (folder / source_path).read_text(encoding='utf-8')
    if artifact_name and (folder / artifact_name).is_file():
        masked = (folder / artifact_name).read_text(encoding='utf-8')
    if not original and not masked: raise HTTPException(409,'该任务没有可恢复的文本')
    return {'task_id':v.task_id,'text':original if v.masked_text is None else v.masked_text.replace(masked,original)}
@app.post('/api/text/restore')
def text_restore(v:RestoreIn): return text_unmask(v)
@app.get('/api/tasks')
def list_tasks():
    c=conn(); rows=c.execute('SELECT id,created,kind,original_name,status,artifact_name,source_sha256,reversible,batch_id,error FROM tasks ORDER BY created DESC').fetchall(); c.close()
    keys=('id','created','kind','filename','status','artifact','sha256','reversible','batch_id','error')
    return {'tasks':[dict(zip(keys,r)) for r in rows]}
@app.delete('/api/tasks/{task_id}')
def delete_task(task_id):
    folder = _task_path(task_id)
    c=conn(); row=c.execute('SELECT id FROM tasks WHERE id=?',(task_id,)).fetchone()
    if not row: c.close(); raise HTTPException(404,'task not found')
    c.execute('DELETE FROM tasks WHERE id=?',(task_id,)); c.commit(); c.close(); shutil.rmtree(folder,ignore_errors=True); return {'deleted':task_id}
@app.get('/api/tasks/{task_id}/artifacts/{name}')
def artifact(task_id,name):
    base=_task_path(task_id); raw_name=str(name); name=Path(raw_name).name
    if raw_name != name or name in {'source.bin','mapping.enc','original.txt'}:
        raise HTTPException(404,'artifact not found')
    c=conn(); row=c.execute('SELECT artifact_name,kind FROM tasks WHERE id=?',(task_id,)).fetchone(); c.close()
    allowed=set(PUBLIC_ARTIFACT_NAMES)
    if row:
        if row[0]: allowed.add(Path(row[0]).name)
        if row[1] == 'image': allowed.update({'masked.png'})
    # Masked/restored files are generated names; permit only the safe suffixes.
    if not (name in allowed or re.fullmatch(r'[A-Za-z0-9._-]+_(?:masked|restored)\.(?:txt|md|docx|xlsx|xlsm|pdf|png|jpg|jpeg|bmp|tif|tiff)', name)):
        raise HTTPException(404,'artifact not found')
    p=(base/name).resolve()
    if p == base or base not in p.parents or not p.is_file(): raise HTTPException(404,'artifact not found')
    return FileResponse(p)
@app.get('/api/rules')
def rules():
    c=conn(); rows=c.execute('SELECT id,name,kind,pattern,entity_type,replacement,enabled FROM rules').fetchall(); c.close(); keys=('id','name','kind','pattern','entity_type','replacement','enabled'); return {'rules':[dict(zip(keys,r)) for r in rows]}
@app.post('/api/rules')
def add_rule(v:RuleIn):
    if v.kind not in {'word','regex'}: raise HTTPException(422,'kind must be word or regex')
    if v.kind=='regex':
        try: re.compile(v.pattern)
        except re.error as e: raise HTTPException(422,f'invalid regex: {e}')
    i=uuid.uuid4().hex; c=conn(); c.execute('INSERT INTO rules VALUES(?,?,?,?,?,?,1)',(i,v.name,v.kind,v.pattern,v.entity_type,v.replacement)); c.commit(); c.close(); return {'id':i}
@app.delete('/api/rules/{rule_id}')
def remove_rule(rule_id):
    c=conn(); c.execute('DELETE FROM rules WHERE id=?',(rule_id,)); c.commit(); c.close(); return {'deleted':rule_id}
@app.post('/api/image/analyze')
async def image_analyze(file:UploadFile=File(...)):
    data=await file.read(MAX_UPLOAD+1)
    if len(data)>MAX_UPLOAD: raise HTTPException(413,'文件超过 50 MB')
    filename = _safe_filename(file.filename or 'original.png'); ext = extension(filename)
    if ext not in IMAGE_EXTENSIONS: raise HTTPException(415,'不支持的图片格式')
    _validate_upload(data, filename, ext)
    try: image=Image.open(io.BytesIO(data)).convert('RGB')
    except Exception as e: raise HTTPException(415,f'无法读取图片: {e}') from e
    i=uuid.uuid4().hex; folder=_task_path(i, create=True); image.save(folder/'original.png'); _register_file_task(i,filename,'image',source_sha256=_sha256(data),status='awaiting_review'); boxes=[]
    ocr_error = None
    if ocr_service.engine:
        try:
            for b in ocr_service.analyze(str(TASKS/i/'original.png')):
                boxes.append({'id':uuid.uuid4().hex,'page':1,'x':b['x1']/image.width,'y':b['y1']/image.height,'width':(b['x2']-b['x1'])/image.width,'height':(b['y2']-b['y1'])/image.height,'text':b['text'],'score':b['score'],'source':'ocr','selected':True,'entity_ids':[]})
        except RuntimeError as exc:
            ocr_error = str(exc)
    payload={'filename':filename,'text':'','entities':[],'boxes':boxes,'image_width':image.width,'image_height':image.height,'sha256':_sha256(data)}
    (folder/'analysis.json').write_text(json.dumps(payload,ensure_ascii=False),encoding='utf-8')
    return {'analysis_id':i,'filename':filename,'image_width':image.width,'image_height':image.height,'ocr_available':ocr_service.engine is not None,'ocr_error':ocr_error,'boxes':boxes,'entities':[],'sha256':payload['sha256']}
@app.post('/api/image/mask')
def image_mask(v:ImageMaskIn):
    src=_task_path(v.analysis_id)/'original.png'
    if not src.is_file(): raise HTTPException(404,'analysis not found')
    _validate_boxes([b.model_dump() for b in v.boxes])
    image=Image.open(src).convert('RGB')
    for b in v.boxes:
        if not b.selected: continue
        l,t=int(b.x*image.width),int(b.y*image.height); r,bt=min(image.width,int((b.x+b.width)*image.width)),min(image.height,int((b.y+b.height)*image.height))
        if r <= l or bt <= t: continue
        image.paste(image.crop((l,t,r,bt)).filter(ImageFilter.GaussianBlur(v.blur_radius)),(l,t))
    image.save(src.parent/'masked.png'); update_task(v.analysis_id,'masked.png',artifact_name='masked.png'); return {'task_id':v.analysis_id,'artifact':f'/api/tasks/{v.analysis_id}/artifacts/masked.png'}
@app.post('/api/document/analyze')
async def document_analyze(file:UploadFile=File(...)):
    data=await file.read(MAX_UPLOAD+1)
    if len(data)>MAX_UPLOAD: raise HTTPException(413,'文件超过 50 MB')
    filename=_safe_filename(file.filename or 'source.bin'); name=filename.lower(); ext=extension(filename); text=''
    if ext not in SUPPORTED: raise HTTPException(415,'不支持的文档格式')
    _validate_upload(data, filename, ext)
    if name.endswith(('.txt','.md')): text=data.decode('utf-8','replace')
    elif name.endswith('.docx'):
        from docx import Document; d=Document(io.BytesIO(data)); text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
    elif name.endswith('.pdf'):
        import fitz; text='\n'.join(p.get_text() for p in fitz.open(stream=data,filetype='pdf'))
    elif not name.endswith(('.png','.jpg','.jpeg','.bmp','.tif','.tiff')): raise HTTPException(415,'不支持的文档格式')
    analysis_id=uuid.uuid4().hex; folder=_task_path(analysis_id, create=True); (folder/'source.bin').write_bytes(data); _register_file_task(analysis_id,filename,'document',source_sha256=_sha256(data),status='awaiting_review')
    entities=analyze(text); (folder/'result.md').write_text(text,encoding='utf-8'); (folder/'content.json').write_text(json.dumps([{'page':1,'type':'text','text':text}],ensure_ascii=False),encoding='utf-8'); (folder/'entities.json').write_text(json.dumps(entities,ensure_ascii=False),encoding='utf-8')
    (folder/'analysis.json').write_text(json.dumps({'filename':filename,'text':text,'entities':entities,'sha256':_sha256(data)},ensure_ascii=False),encoding='utf-8')
    return {'analysis_id':analysis_id,'filename':filename,'text':text,'entities':entities,'content':[{'page':1,'type':'text','text':text}],'sha256':_sha256(data),'artifacts':[f'/api/tasks/{analysis_id}/artifacts/result.md',f'/api/tasks/{analysis_id}/artifacts/content.json',f'/api/tasks/{analysis_id}/artifacts/entities.json']}
@app.post('/api/document/mask')
def document_mask(v:DocumentMaskIn):
    if not v.analysis_id:
        result = text_mask(v)
        c=conn(); c.execute('UPDATE tasks SET kind=? WHERE id=?',('document',result['task_id'])); c.commit(); c.close()
        return result
    folder=_task_path(v.analysis_id)
    if not (folder/'source.bin').is_file(): raise HTTPException(404,'analysis not found')
    masked=apply_mask(v.text,v.entities)
    (folder/'masked.md').write_text(masked,encoding='utf-8'); update_task(v.analysis_id,'masked.md')
    return {'task_id':v.analysis_id,'masked_text':masked,'artifact':f'/api/tasks/{v.analysis_id}/artifacts/masked.md'}

# Unified final-version file and batch APIs. The legacy routes above remain available.
class FileMaskIn(BaseModel):
    analysis_id: str
    filename: str
    text: str = ''
    entities: list[dict[str, Any]] = Field(default_factory=list)
    policies: dict[str, dict[str, Any]] = Field(default_factory=dict)
    reviewed: bool = False
    reversible: bool = False
    password: str | None = None
    boxes: list[dict[str, Any]] = Field(default_factory=list)

class BatchMaskIn(BaseModel):
    policies: dict[str, dict[str, Any]] = Field(default_factory=dict)
    reviewed: bool = False
    reversible: bool = False
    password: str | None = None
    items: list[dict[str, Any]] = Field(default_factory=list)

def _task_folder(task_id: str) -> Path:
    return _task_path(task_id, create=True)

def _register_file_task(task_id: str, filename: str, kind: str, *, source_sha256='',
                        source_path='source.bin', status='analyzing', batch_id='', metadata=None) -> None:
    register_task(task_id, kind, original_name=_safe_filename(filename), source_path=source_path,
                  source_sha256=source_sha256, status=status, batch_id=batch_id, metadata=metadata)

def _analysis_payload(task_id: str) -> dict[str, Any]:
    folder = _task_path(task_id)
    path = folder / 'analysis.json'
    if not path.is_file():
        raise HTTPException(404, 'analysis not found')
    try:
        payload = json.loads(path.read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError) as exc:
        raise HTTPException(409, '分析结果已损坏') from exc
    if not isinstance(payload, dict):
        raise HTTPException(409, '分析结果格式无效')
    return payload

def _entity_key(entity: dict[str, Any]) -> tuple[Any, ...]:
    bbox = entity.get('bbox') or {}
    return (str(entity.get('type', '')), str(entity.get('text', '')), int(entity.get('start', 0)),
            int(entity.get('end', 0)), round(float(bbox.get('x', 0)), 8), round(float(bbox.get('y', 0)), 8),
            round(float(bbox.get('width', 0)), 8), round(float(bbox.get('height', 0)), 8), int(entity.get('page', 1)))

def _review_entities(server_entities: list[dict[str, Any]], submitted: list[dict[str, Any]], *, allow_new=False) -> list[dict[str, Any]]:
    """Accept only selection/box edits over the server analysis snapshot."""
    by_id = {str(item.get('id')): item for item in submitted if item.get('id')}
    by_key = {_entity_key(item): item for item in submitted}
    result=[]
    for original in server_entities:
        candidate = by_id.get(str(original.get('id'))) or by_key.get(_entity_key(original))
        merged = dict(original)
        if candidate is not None:
            immutable = ('type', 'text', 'start', 'end')
            if any(str(candidate.get(k, '')) != str(original.get(k, '')) for k in immutable):
                raise HTTPException(422, '实体内容与分析结果不一致')
            merged['selected'] = bool(candidate.get('selected', True))
            if original.get('bbox') and candidate.get('bbox'):
                merged['bbox'] = candidate['bbox']
        else:
            merged['selected'] = False
        result.append(merged)
    if not allow_new:
        known_ids = {str(item.get('id')) for item in server_entities}
        known_keys = {_entity_key(item) for item in server_entities}
        for candidate in submitted:
            if str(candidate.get('id')) not in known_ids and _entity_key(candidate) not in known_keys:
                raise HTTPException(422, '提交了未来源于分析结果的实体')
    else:
        for candidate in submitted:
            if str(candidate.get('id')) not in {str(x.get('id')) for x in server_entities} and _entity_key(candidate) not in {_entity_key(x) for x in server_entities}:
                if candidate.get('selected', True): result.append(dict(candidate))
    return result

def _validate_boxes(boxes: list[dict[str, Any]]) -> list[dict[str, Any]]:
    validated=[]
    for box in boxes:
        try:
            x,y,w,h = (float(box.get(k, 0)) for k in ('x','y','width','height'))
        except (TypeError, ValueError) as exc:
            raise HTTPException(422, '图片框坐标无效') from exc
        if not all(map(lambda value: value == value and abs(value) != float('inf'), (x,y,w,h))) or w <= 0 or h <= 0 or x < 0 or y < 0 or x >= 1 or y >= 1:
            raise HTTPException(422, '图片框坐标超出范围')
        if x + w <= 0 or y + h <= 0:
            raise HTTPException(422, '图片框尺寸无效')
        item=dict(box); item.update({'x':max(0.0,x),'y':max(0.0,y),'width':min(w,1-x),'height':min(h,1-y)})
        validated.append(item)
    return validated

@app.get('/api/policies')
def get_policies():
    return {'policies': policy_store}

@app.put('/api/policies')
def put_policies(value: dict[str, dict[str, Any]]):
    for entity_type, policy in value.items():
        if policy.get('text_action', 'token') not in {'keep', 'token', 'replace', 'mask', 'hash', 'pseudonym'}:
            raise HTTPException(422, f'不支持的文本策略: {entity_type}')
        if policy.get('image_action', 'solid') not in {'keep', 'blur', 'pixelate', 'solid', 'text'}:
            raise HTTPException(422, f'不支持的图片策略: {entity_type}')
    policy_store.update(value)
    return {'policies': policy_store}

@app.post('/api/files/analyze')
async def files_analyze(file: UploadFile = File(...)):
    data = await file.read(MAX_UPLOAD + 1)
    if len(data) > MAX_UPLOAD:
        raise HTTPException(413, '文件超过 50 MB')
    filename = _safe_filename(file.filename or 'upload.bin')
    ext = extension(filename)
    if ext not in SUPPORTED:
        raise HTTPException(415, f'不支持的文件格式: .{ext}')
    _validate_upload(data, filename, ext)
    task_id = uuid.uuid4().hex
    folder = _task_folder(task_id)
    try:
        (folder / 'source.bin').write_bytes(data)
        text = extract_text(data, filename)
    except HTTPException:
        shutil.rmtree(folder, ignore_errors=True)
        raise
    except Exception as exc:
        shutil.rmtree(folder, ignore_errors=True)
        raise HTTPException(415, f'文件解析失败: {exc}') from exc
    entities = analyze(text) if text else []
    warnings=[]
    if ext in {'png', 'jpg', 'jpeg', 'bmp', 'tif', 'tiff'} and ocr_service.engine:
        try:
            image = Image.open(io.BytesIO(data)).convert('RGB')
            for box in ocr_service.analyze(str(folder / 'source.bin')):
                box_entity = analyze(box['text'])
                for entity in box_entity:
                    entity['bbox'] = {'x': box['x1'] / image.width, 'y': box['y1'] / image.height, 'width': (box['x2'] - box['x1']) / image.width, 'height': (box['y2'] - box['y1']) / image.height}
                    entity['page'] = 1
                    entities.append(entity)
        except Exception as exc:
            warnings.append({'code': 'OCR_FAILED', 'message': str(exc)})
    manifest = content_manifest(filename, text, entities) | {'sha256': _sha256(data), 'kind': ext, 'warnings': warnings}
    _register_file_task(task_id, filename, ext, source_sha256=manifest['sha256'], status='awaiting_review', metadata={'manifest_version': 1})
    (folder / 'analysis.json').write_text(json.dumps(manifest, ensure_ascii=False), encoding='utf-8')
    return {'analysis_id': task_id, 'filename': filename, 'kind': ext, 'text': text, 'entities': entities, 'warnings': warnings, 'sha256': manifest['sha256'], 'artifacts': [f'/api/tasks/{task_id}/artifacts/analysis.json']}

@app.post('/api/files/mask')
def files_mask(value: FileMaskIn):
    if not value.reviewed:
        raise HTTPException(409, '请先完成实体复核')
    folder = _task_folder(value.analysis_id)
    source = folder / 'source.bin'
    if not source.is_file():
        raise HTTPException(404, 'analysis not found')
    if value.reversible and not value.password:
        raise HTTPException(422, '可逆模式需要口令')
    c = conn(); row = c.execute('SELECT original_name,original,source_sha256 FROM tasks WHERE id=?', (value.analysis_id,)).fetchone(); c.close()
    if not row:
        raise HTTPException(404, 'analysis not found')
    server_filename = str(row[0] or row[1])
    if extension(value.filename) != extension(server_filename):
        raise HTTPException(422, '文件扩展名与分析任务不一致')
    data = source.read_bytes()
    expected_hash = str(row[2] or '')
    if expected_hash and _sha256(data) != expected_hash:
        raise HTTPException(409, '分析源文件已被修改')
    manifest = _analysis_payload(value.analysis_id)
    server_text = str(manifest.get('text', ''))
    # Text may be normalized by clients; the signed source and reviewed entity snapshot are authoritative.
    server_entities = manifest.get('entities', [])
    if not isinstance(server_entities, list):
        raise HTTPException(409, '分析实体格式无效')
    allow_new = extension(server_filename) in IMAGE_EXTENSIONS
    entities = _review_entities(server_entities, value.entities, allow_new=allow_new)
    boxes = _validate_boxes(value.boxes)
    try:
        masked = mask_file(data, server_filename, server_text, entities, value.policies or policy_store, boxes)
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    except (OSError, RuntimeError, zipfile.BadZipFile) as exc:
        raise HTTPException(422, f'文件脱敏失败: {exc}') from exc
    ext = extension(server_filename)
    output_name = f'{Path(server_filename).stem}_masked.{ext}'
    (folder / output_name).write_bytes(masked)
    report = content_manifest(server_filename, server_text, entities) | {'policies': value.policies or policy_store, 'reversible': value.reversible, 'sha256': expected_hash}
    (folder / 'report.json').write_text(json.dumps(report, ensure_ascii=False), encoding='utf-8')
    if value.reversible:
        try:
            (folder / 'mapping.enc').write_bytes(encrypt_json({'version': 1, 'filename': server_filename, 'data': data.hex(), 'sha256': expected_hash}, value.password or ''))
        except ValueError as exc:
            raise HTTPException(422, str(exc)) from exc
    # Keep plaintext source only until a successful result has been produced.
    source.unlink(missing_ok=True)
    update_task(value.analysis_id, output_name, artifact_name=output_name, source_path='', reversible=value.reversible)
    return {'task_id': value.analysis_id, 'status': 'completed', 'artifact': f'/api/tasks/{value.analysis_id}/artifacts/{output_name}', 'report': f'/api/tasks/{value.analysis_id}/artifacts/report.json', 'reversible': value.reversible}

@app.post('/api/tasks/{task_id}/restore')
def restore_file(task_id: str, value: RestoreFileIn):
    folder = _task_folder(task_id)
    mapping = folder / 'mapping.enc'
    if not mapping.is_file():
        raise HTTPException(409, '该任务未启用可逆模式')
    try:
        payload = decrypt_json(mapping.read_bytes(), value.password)
    except Exception as exc:
        raise HTTPException(403, '口令错误或映射已损坏') from exc
    filename = _safe_filename(str(payload.get('filename', 'restored.bin')))
    try:
        restored = bytes.fromhex(str(payload['data']))
    except (KeyError, TypeError, ValueError) as exc:
        raise HTTPException(409, '加密映射格式无效') from exc
    if payload.get('sha256') and _sha256(restored) != payload['sha256']:
        raise HTTPException(409, '恢复文件校验失败')
    restored_name = f'{Path(filename).stem}_restored.{extension(filename)}'
    (folder / restored_name).write_bytes(restored)
    update_task(task_id, restored_name, artifact_name=restored_name)
    return {'task_id': task_id, 'artifact': f'/api/tasks/{task_id}/artifacts/{restored_name}'}

batch_lock = threading.RLock()

def _json_state(state: dict[str, Any]) -> dict[str, Any]:
    """Return a JSON-safe copy without upload bytes or passwords."""
    result = {}
    for key, value in state.items():
        if key in {'password', 'files_bytes'}:
            continue
        result[key] = value
    return result

def _save_batch(batch_id: str, state: dict[str, Any]) -> None:
    now = datetime.now().timestamp(); state['updated'] = now
    payload = _json_state(state)
    c=conn(); c.execute('''INSERT INTO batches(id,created,updated,status,total,completed,failed,archive_name,metadata_json,error)
        VALUES(?,?,?,?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET updated=excluded.updated,status=excluded.status,
        total=excluded.total,completed=excluded.completed,failed=excluded.failed,archive_name=excluded.archive_name,
        metadata_json=excluded.metadata_json,error=excluded.error''',
        (batch_id, state.get('created', now), now, state.get('status','queued'), state.get('total',0),
         state.get('completed',0), state.get('failed',0), state.get('archive_name',''),
         json.dumps(payload, ensure_ascii=False), state.get('error','')))
    c.commit(); c.close()

def _public_batch(state: dict[str, Any]) -> dict[str, Any]:
    result = deepcopy(_json_state(state))
    # Internal input paths are never exposed to the browser.
    for item in result.get('files', []):
        item.pop('storage_name', None)
    return result

def _read_batch_input(batch_id: str, metadata: dict[str, Any]) -> bytes:
    path = _task_path(batch_id) / 'inputs' / str(metadata.get('storage_name', ''))
    if path.parent.name != 'inputs' or not path.is_file():
        raise FileNotFoundError('batch input not found')
    data = path.read_bytes()
    expected = str(metadata.get('sha256', ''))
    if expected and _sha256(data) != expected:
        raise ValueError('批次源文件校验失败')
    return data

def _analyze_batch_item(batch_id: str, metadata: dict[str, Any]) -> dict[str, Any]:
    filename = str(metadata['filename']); data = _read_batch_input(batch_id, metadata)
    task_id = uuid.uuid4().hex; folder = _task_path(task_id, create=True); (folder / 'source.bin').write_bytes(data)
    text = extract_text(data, filename) if extension(filename) not in IMAGE_EXTENSIONS else ''
    entities = analyze(text) if text else []
    manifest = content_manifest(filename, text, entities) | {'kind': extension(filename), 'sha256': _sha256(data), 'warnings': []}
    (folder / 'analysis.json').write_text(json.dumps(manifest, ensure_ascii=False), encoding='utf-8')
    _register_file_task(task_id, filename, extension(filename), source_sha256=manifest['sha256'], status='awaiting_review', batch_id=batch_id)
    return {'filename': filename, 'task_id': task_id, 'status': 'awaiting_review', 'text': text, 'entities': entities, 'boxes': [], 'sha256': manifest['sha256']}

def _unique_archive_name(directory: Path, filename: str, used: set[str]) -> str:
    base = re.sub(r'[^A-Za-z0-9._-]+', '_', Path(filename).name)[:180] or 'result.bin'
    stem, suffix = Path(base).stem, Path(base).suffix
    candidate = f'{stem}_masked{suffix}'
    index = 2
    while candidate in used or (directory / candidate).exists():
        candidate = f'{stem}_masked_{index}{suffix}'; index += 1
    used.add(candidate); return candidate

def _run_batch(batch_id: str, options: BatchMaskIn | None = None):
    with batch_lock:
        state = batches.get(batch_id)
    if not state:
        return
    # First phase always analyzes and waits for an explicit review request.
    if not state.get('analysis_ready'):
        state['status'] = 'analyzing'; state['items'] = []; _save_batch(batch_id, state)
        for metadata in state.get('files', []):
            if state.get('cancelled'): break
            try:
                state['items'].append(_analyze_batch_item(batch_id, metadata))
            except (OSError, ValueError, RuntimeError, HTTPException) as exc:
                state['items'].append({'filename': metadata.get('filename',''), 'status': 'failed', 'error': str(exc)})
                state['failed'] = int(state.get('failed', 0)) + 1
            _save_batch(batch_id, state)
        state['analysis_ready'] = True
        state['status'] = 'cancelled' if state.get('cancelled') else 'awaiting_review'
        _save_batch(batch_id, state)
        return
    if not options or not options.reviewed:
        state['status'] = 'awaiting_review'; _save_batch(batch_id, state); return
    state['status'] = 'running'; state['completed'] = 0; state['failed'] = 0
    archive_dir = _task_path(batch_id, create=True); archive_dir.mkdir(parents=True, exist_ok=True); used=set()
    submitted = {str(item.get('task_id')): item for item in options.items if item.get('task_id')}
    for item in state.get('items', []):
        if state.get('cancelled'): break
        if item.get('status') == 'failed': continue
        task_id = str(item.get('task_id')); task_folder = _task_path(task_id); source = task_folder / 'source.bin'
        current = submitted.get(task_id, {})
        item['status'] = 'running'; _save_batch(batch_id, state)
        try:
            manifest = _analysis_payload(task_id); text = str(manifest.get('text',''))
            if current.get('text', text) != text:
                raise ValueError('批次复核文本与分析结果不一致')
            entities = _review_entities(manifest.get('entities', []), current.get('entities', manifest.get('entities', [])), allow_new=extension(item['filename']) in IMAGE_EXTENSIONS)
            boxes = _validate_boxes(current.get('boxes', []))
            data = source.read_bytes(); expected = str(manifest.get('sha256',''))
            if expected and _sha256(data) != expected: raise ValueError('分析源文件已被修改')
            masked = mask_file(data, item['filename'], text, entities, options.policies or policy_store, boxes)
            output_name = f'{Path(item["filename"]).stem}_masked.{extension(item["filename"])}'
            (task_folder / output_name).write_bytes(masked)
            if options.reversible:
                if not options.password: raise ValueError('可逆模式需要口令')
                (task_folder / 'mapping.enc').write_bytes(encrypt_json({'version':1,'filename':item['filename'],'data':data.hex(),'sha256':expected}, options.password))
            source.unlink(missing_ok=True)
            archive_name = _unique_archive_name(archive_dir, item['filename'], used); (archive_dir / archive_name).write_bytes(masked)
            update_task(task_id, output_name, artifact_name=output_name, source_path='', reversible=options.reversible)
            item.update({'status':'completed','artifact':archive_name}); state['completed'] += 1
        except Exception as exc:
            item.update({'status':'failed','error':str(exc)}); state['failed'] += 1
        _save_batch(batch_id, state)
    final_status = 'cancelled' if state.get('cancelled') else ('partial' if state.get('failed') else 'completed')
    report = {'batch_id':batch_id,'status':final_status,'items':state.get('items',[]),'completed':state.get('completed',0),'failed':state.get('failed',0)}
    (archive_dir / 'report.json').write_text(json.dumps(report,ensure_ascii=False),encoding='utf-8')
    with zipfile.ZipFile(archive_dir / 'results.zip','w',zipfile.ZIP_DEFLATED) as archive:
        for path in archive_dir.iterdir():
            if path.name != 'results.zip' and path.is_file(): archive.write(path,path.name)
    state['status'] = final_status; state['archive'] = f'/api/tasks/{batch_id}/artifacts/results.zip'; state['archive_name']='results.zip'; state['report']=report
    _save_batch(batch_id, state)

@app.post('/api/batches')
async def create_batch(files: list[UploadFile] = File(...), options: str = Form('{}')):
    if len(files) > MAX_BATCH_FILES:
        raise HTTPException(413, f'每批最多 {MAX_BATCH_FILES} 个文件')
    try:
        payload = BatchMaskIn.model_validate_json(options)
    except ValidationError as exc:
        raise HTTPException(422, f'批次参数无效: {exc}') from exc
    # Upload always enters review, even if a client sends reviewed=true.
    if payload.reversible and payload.password and len(payload.password) < 12:
        raise HTTPException(422, '可逆模式口令至少需要 12 个字符')
    items: list[dict[str, Any]] = []; total = 0; batch_id = uuid.uuid4().hex; batch_dir = _task_path(batch_id, create=True); input_dir = batch_dir / 'inputs'; input_dir.mkdir(exist_ok=True)
    for upload in files:
        filename = _safe_filename(upload.filename or 'upload.bin'); ext = extension(filename); data = await upload.read(MAX_UPLOAD + 1); total += len(data)
        if len(data) > MAX_UPLOAD: raise HTTPException(413, f'{upload.filename} 超过 50 MB')
        if ext not in SUPPORTED: raise HTTPException(415, f'不支持的文件格式: {upload.filename}')
        _validate_upload(data, filename, ext)
        storage_name = f'{len(items):04d}_{uuid.uuid4().hex}{Path(filename).suffix.lower()}'
        (input_dir / storage_name).write_bytes(data)
        items.append({'filename':filename,'storage_name':storage_name,'sha256':_sha256(data),'kind':ext})
    if total > MAX_BATCH_BYTES: raise HTTPException(413, f'批次超过 {MAX_BATCH_BYTES} 字节')
    now=datetime.now().timestamp(); batches[batch_id] = {'status':'queued','created':now,'updated':now,'total':len(items),'completed':0,'failed':0,'files':items,'policies':payload.policies or {},'reversible':payload.reversible,'cancelled':False}
    _save_batch(batch_id,batches[batch_id])
    threading.Thread(target=_run_batch, args=(batch_id, None), daemon=True).start()
    return {'batch_id': batch_id, 'status': 'queued'}

@app.get('/api/batches/{batch_id}')
def batch_status(batch_id: str):
    if not _valid_task_id(batch_id): raise HTTPException(404,'batch not found')
    state=batches.get(batch_id)
    if state is None:
        c=conn(); row=c.execute('SELECT metadata_json,status FROM batches WHERE id=?',(batch_id,)).fetchone(); c.close()
        if not row: raise HTTPException(404,'batch not found')
        try: state=json.loads(row[0]); state['status']=row[1]
        except json.JSONDecodeError: state={'status':'failed','error':'批次状态损坏'}
    return _public_batch(state)

@app.post('/api/batches/{batch_id}/cancel')
def cancel_batch(batch_id: str):
    if not _valid_task_id(batch_id) or batch_id not in batches: raise HTTPException(404, 'batch not found')
    state=batches[batch_id]; state['cancelled']=True; state['status']='cancelled'; _save_batch(batch_id,state); return {'batch_id': batch_id, 'status': 'cancelled'}

@app.post('/api/batches/{batch_id}/mask')
def mask_batch(batch_id: str, value: BatchMaskIn):
    state = batches.get(batch_id)
    if not state:
        raise HTTPException(404, 'batch not found')
    if state.get('status') != 'awaiting_review':
        raise HTTPException(409, '批次当前不在待复核状态')
    if not value.reviewed:
        raise HTTPException(409, '请先完成实体复核')
    if value.reversible and not value.password:
        raise HTTPException(422, '可逆模式需要口令')
    if value.reversible and len(value.password or '') < 12:
        raise HTTPException(422, '可逆模式口令至少需要 12 个字符')
    state['status'] = 'queued'
    state['reversible'] = value.reversible; state['policies'] = value.policies or policy_store; state['cancelled'] = False
    state['reviewed_items'] = value.items
    _save_batch(batch_id,state)
    threading.Thread(target=_run_batch, args=(batch_id, value), daemon=True).start()
    return {'batch_id': batch_id, 'status': 'queued'}

@app.get('/api/batches/{batch_id}/download')
def download_batch(batch_id: str):
    if not _valid_task_id(batch_id): raise HTTPException(404,'batch not found')
    state = batches.get(batch_id)
    if not state or not state.get('archive'): raise HTTPException(404, 'batch result not ready')
    archive = _task_path(batch_id) / 'results.zip'
    if not archive.is_file(): raise HTTPException(404,'batch result not found')
    return FileResponse(archive, filename=f'{batch_id}_results.zip', media_type='application/zip')
app.mount('/static',StaticFiles(directory=ROOT/'app'/'static'),name='static')
@app.get('/')
def index(): return FileResponse(ROOT/'app'/'static'/'index.html')

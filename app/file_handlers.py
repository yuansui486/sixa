from __future__ import annotations

import io
import math
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw, ImageFilter, ImageFont, ImageOps

from .policies import apply_entities, policy_for, replacement_for

SUPPORTED = {"txt", "md", "docx", "xlsx", "xlsm", "pdf", "png", "jpg", "jpeg", "bmp", "tif", "tiff"}


def extension(filename: str) -> str:
    return Path(filename or "").suffix.lower().lstrip(".")


def _text_from_docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    return "\n".join(value for _kind, _owner, value in _docx_text_fields(document))


def _text_from_xlsx(data: bytes, keep_vba: bool = False) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=False, data_only=False, keep_vba=keep_vba)
    try:
        return "\n".join(value for _kind, _owner, value in _xlsx_fields(workbook))
    finally:
        vba_archive = getattr(workbook, "vba_archive", None)
        if vba_archive is not None:
            vba_archive.close()


def extract_text(data: bytes, filename: str) -> str:
    ext = extension(filename)
    if ext in {"txt", "md"}:
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        return data.decode("utf-8", "replace")
    if ext == "docx":
        return _text_from_docx(data)
    if ext in {"xlsx", "xlsm"}:
        return _text_from_xlsx(data, keep_vba=ext == "xlsm")
    if ext == "pdf":
        import fitz

        document = fitz.open(stream=data, filetype="pdf")
        return "\n".join(page.get_text() for page in document)
    return ""


def _docx_parts(document: Any) -> list[tuple[list[Any], str]]:
    """Return paragraph groups in the same order as ``_text_from_docx``.

    The text analyzer sees one newline-separated stream.  Keeping the exact
    source ranges here lets us apply an explicitly reviewed entity to only the
    occurrence selected by the user, rather than replacing every equal string
    in the document.
    """
    parts: list[tuple[list[Any], str]] = []
    for paragraph in document.paragraphs:
        parts.append(([paragraph], paragraph.text))
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs = list(cell.paragraphs)
                parts.append((paragraphs, "\n".join(p.text for p in paragraphs)))
    for section in document.sections:
        header = section.header
        footer = section.footer
        for paragraph in header.paragraphs:
            parts.append(([paragraph], paragraph.text))
        for table in header.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs = list(cell.paragraphs)
                    parts.append((paragraphs, "\n".join(p.text for p in paragraphs)))
        for paragraph in footer.paragraphs:
            parts.append(([paragraph], paragraph.text))
        for table in footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs = list(cell.paragraphs)
                    parts.append((paragraphs, "\n".join(p.text for p in paragraphs)))
    # Comments are stored in word/comments.xml and are not represented by
    # ``document.paragraphs``.  They still ship with the sanitized DOCX, so
    # expose them to the same review stream as visible body text.
    comments = getattr(document, "comments", None)
    if comments is not None:
        for comment in comments:
            for paragraph in comment.paragraphs:
                parts.append(([paragraph], paragraph.text))
            for table in comment.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs = list(cell.paragraphs)
                        parts.append((paragraphs, "\n".join(p.text for p in paragraphs)))
    return parts


def _paragraph_hyperlink_relationships(paragraph: Any) -> list[Any]:
    """Return external hyperlink relationships referenced by a paragraph."""
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.oxml.ns import qn

    relationships: list[Any] = []
    for element in paragraph._p.iter(qn("w:hyperlink")):
        relationship_id = element.get(qn("r:id"))
        if not relationship_id:
            continue
        relationship = paragraph.part.rels.get(relationship_id)
        if (
            relationship is not None
            and relationship.is_external
            and relationship.reltype == RELATIONSHIP_TYPE.HYPERLINK
        ):
            relationships.append(relationship)
    return relationships


def _docx_text_fields(document: Any) -> list[tuple[str, Any, str]]:
    """Build the exact review stream for DOCX text and hidden link targets."""
    from docx.opc.constants import RELATIONSHIP_TYPE

    fields: list[tuple[str, Any, str]] = []
    referenced_relationships: set[int] = set()
    for paragraphs, _part_text in _docx_parts(document):
        for paragraph in paragraphs:
            # Put the target immediately before its display paragraph.  This
            # preserves the historical ordering of visible duplicate values
            # while making the hidden relationship target independently
            # reviewable.
            for relationship in _paragraph_hyperlink_relationships(paragraph):
                fields.append(("hyperlink", relationship, str(relationship.target_ref)))
                referenced_relationships.add(id(relationship))
            fields.append(("paragraph", paragraph, paragraph.text))

    # Include valid external hyperlink relationships that are not referenced
    # by a standard w:hyperlink node (for example links in less common parts).
    package_parts = sorted(
        document.part.package.parts,
        key=lambda part: str(getattr(part, "partname", "")),
    )
    for part in package_parts:
        relationships = getattr(part, "rels", None)
        if relationships is None:
            continue
        for relationship_id in sorted(relationships):
            relationship = relationships[relationship_id]
            if (
                id(relationship) not in referenced_relationships
                and relationship.is_external
                and relationship.reltype == RELATIONSHIP_TYPE.HYPERLINK
            ):
                fields.append(("hyperlink", relationship, str(relationship.target_ref)))
                referenced_relationships.add(id(relationship))
    return fields


def _docx_field_ranges(document: Any) -> list[tuple[str, Any, str, int, int]]:
    fields = _docx_text_fields(document)
    ranges: list[tuple[str, Any, str, int, int]] = []
    cursor = 0
    for index, (kind, owner, value) in enumerate(fields):
        ranges.append((kind, owner, value, cursor, cursor + len(value)))
        cursor += len(value)
        if index + 1 < len(fields):
            cursor += 1
    return ranges


def _docx_paragraph_ranges(document: Any) -> list[tuple[Any, int, int]]:
    return [
        (owner, start, end)
        for kind, owner, _value, start, end in _docx_field_ranges(document)
        if kind == "paragraph"
    ]


def _clear_docx_metadata(document: Any) -> None:
    """Remove identifying package and comment metadata from a DOCX output."""
    properties = document.core_properties
    for attribute in (
        "author", "category", "comments", "content_status", "identifier",
        "keywords", "language", "last_modified_by", "subject", "title", "version",
    ):
        try:
            setattr(properties, attribute, "")
        except (AttributeError, TypeError, ValueError):
            pass
    for attribute in ("created", "last_printed", "modified"):
        try:
            setattr(properties, attribute, None)
        except (AttributeError, TypeError, ValueError):
            pass
    comments = getattr(document, "comments", None)
    if comments is None:
        return
    from docx.oxml.ns import qn

    for comment in comments:
        comment.author = ""
        comment.initials = ""
        # python-docx exposes the timestamp read-only, but it is metadata and
        # can be removed safely from the underlying OOXML element.
        comment._element.attrib.pop(qn("w:date"), None)


def _replace_docx_text(data: bytes, entities: list[dict], policies: dict | None) -> bytes:
    from docx import Document

    document = Document(io.BytesIO(data))
    mapping: dict[tuple[str, str], str] = {}
    ranges = _docx_field_ranges(document)
    # A header/footer can be linked by several sections.  The same XML
    # paragraph then appears at several positions in the extracted stream;
    # collect all reviewed occurrences before mutating it once.
    grouped: dict[int, tuple[Any, list[dict]]] = {}
    hyperlink_groups: dict[int, tuple[Any, str, list[dict]]] = {}
    for kind, owner, value, global_start, global_end in ranges:
        if kind == "paragraph":
            identity = id(owner._p)
            item = grouped.setdefault(identity, (owner, []))
        else:
            identity = id(owner)
            item = hyperlink_groups.setdefault(identity, (owner, value, []))
        local_entities = item[-1]
        for entity in entities:
            target = str(entity.get("text", ""))
            if not target:
                continue
            has_offsets = "start" in entity and "end" in entity
            if has_offsets:
                try:
                    start, end = int(entity["start"]), int(entity["end"])
                except (TypeError, ValueError):
                    continue
                if global_start <= start and end <= global_end and end > start:
                    local = dict(entity)
                    local["start"], local["end"] = start - global_start, end - global_start
                    local_entities.append(local)
            elif target in value:
                # Compatibility for direct callers that provide only text.
                local_entities.extend(
                    {**entity, "start": match.start(), "end": match.end()}
                    for match in re.finditer(re.escape(target), value)
                )
    for paragraph, local_entities in grouped.values():
        if local_entities:
            # Repeated linked headers may map the same local span more than
            # once.  Keep one copy so replacement mapping and overlap logic
            # remain deterministic.
            unique: dict[tuple[int, int, str], dict] = {
                (int(item["start"]), int(item["end"]), str(item.get("type", ""))): item
                for item in local_entities
            }
            _replace_docx_paragraph(paragraph, list(unique.values()), policies, mapping)
    for relationship, target, local_entities in hyperlink_groups.values():
        if not local_entities:
            continue
        unique = {
            (int(item["start"]), int(item["end"]), str(item.get("type", ""))): item
            for item in local_entities
        }
        masked_target, mapping = apply_entities(target, list(unique.values()), policies, mapping)
        # python-docx exposes target_ref as read-only.  Updating the underlying
        # relationship target is the supported serialization path used by its
        # relationship writer and keeps the rest of the package unchanged.
        relationship._target = masked_target
    _clear_docx_metadata(document)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _run_ranges(original: str, runs: list[Any]) -> list[tuple[int, int]]:
    ranges: list[tuple[int, int]] = []
    cursor = 0
    for run in runs:
        end = cursor + len(run.text or "")
        ranges.append((cursor, end))
        cursor = end
    return ranges


def _replace_docx_paragraph(
    paragraph: Any,
    entities: list[dict],
    policies: dict | None,
    mapping: dict[tuple[str, str], str] | None = None,
) -> None:
    # ``Paragraph.runs`` contains only direct children.  Hyperlinks and some
    # field constructs keep their runs nested in another element, so inspect
    # every ``w:r`` in document order and mutate those runs in place.
    from docx.oxml.ns import qn
    from docx.text.run import Run

    run_xmls = list(paragraph._p.iter(qn("w:r")))
    runs = [Run(run_xml, paragraph) for run_xml in run_xmls]
    if not runs:
        return
    original = "".join(run.text or "" for run in runs)
    if not original:
        return
    mapping = mapping if mapping is not None else {}
    spans: list[dict] = []
    for entity in entities:
        try:
            start, end = int(entity["start"]), int(entity["end"])
        except (KeyError, TypeError, ValueError):
            continue
        target = str(entity.get("text", ""))
        if not target or start < 0 or end <= start or end > len(original):
            continue
        if original[start:end] != target:
            raise ValueError("实体区间与 DOCX 段落文本不一致")
        if entity.get("selected", True):
            spans.append({**entity, "start": start, "end": end})
    if not spans:
        return
    spans.sort(key=lambda item: (int(item["start"]), int(item["end"])), reverse=True)
    non_overlapping: list[dict] = []
    cursor = len(original) + 1
    for span in spans:
        if int(span["end"]) > cursor:
            continue
        non_overlapping.append(span)
        cursor = int(span["start"])

    # Process from right to left.  Replacements can change the length of the
    # text to their right, but never invalidate offsets for an earlier span.
    for span in non_overlapping:
        start, end = int(span["start"]), int(span["end"])
        current = "".join(run.text or "" for run in runs)
        if start < 0 or end > len(current) or current[start:end] != span.get("text"):
            raise ValueError("实体区间与 DOCX 段落文本不一致")
        replacement = replacement_for(span, policies, mapping)
        ranges = _run_ranges(current, runs)
        affected = [
            (index, run, run_start, run_end)
            for index, (run_start, run_end) in enumerate(ranges)
            for run in [runs[index]]
            if run_start < end and run_end > start
        ]
        if not affected:
            continue
        first_index, first_run, first_start, _first_end = affected[0]
        last_index, last_run, last_start, _last_end = affected[-1]
        if first_index == last_index:
            local_start, local_end = start - first_start, end - first_start
            first_run.text = (first_run.text or "")[:local_start] + replacement + (first_run.text or "")[local_end:]
            continue
        first_local_start = start - first_start
        first_run.text = (first_run.text or "")[:first_local_start] + replacement
        for _index, run, run_start, run_end in affected[1:-1]:
            run.text = ""
        last_local_end = end - last_start
        last_run.text = (last_run.text or "")[last_local_end:]


def _xlsx_fields(workbook: Any) -> list[tuple[str, Any | None, str]]:
    """Yield fields in the same order as :func:`_text_from_xlsx`."""
    fields: list[tuple[str, Any | None, str]] = []
    for sheet in workbook.worksheets:
        # Brackets keep sheet boundaries easy to scan in the review text;
        # retaining the owner lets a selected title be changed in the output.
        fields.append(("sheet", sheet, f"[{sheet.title}]"))
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    value = str(cell.value)
                    # Keep formulas as a distinct field.  Formula literals are
                    # still analyzed and masked, but require a syntax-aware
                    # replacement so a token such as ``138****8000`` cannot
                    # turn an unquoted numeric operand into an invalid formula.
                    fields.append(("formula" if value.startswith("=") else "value", cell, value))
                if cell.comment:
                    fields.append(("comment", cell, str(cell.comment.text or "")))
                if cell.hyperlink and cell.hyperlink.target:
                    fields.append(("url", cell, str(cell.hyperlink.target)))
        # Page headers/footers are serialized independently of cell values and
        # are frequently used for printed customer reports.  Include all six
        # OOXML variants in the review stream.
        for attribute in (
            "oddHeader", "evenHeader", "firstHeader",
            "oddFooter", "evenFooter", "firstFooter",
        ):
            item = getattr(sheet, attribute, None)
            if item is None:
                continue
            for position in ("left", "center", "right"):
                text = getattr(getattr(item, position, None), "text", None)
                if text:
                    fields.append(("header_footer", (item, position), str(text)))
    # Defined names can themselves carry identifiers.  Keep their actual
    # serialized name/reference rather than the object's diagnostic repr.
    for name, defined in workbook.defined_names.items():
        fields.append(("defined_name", defined, str(name)))
        fields.append(("defined_value", defined, str(getattr(defined, "attr_text", "") or "")))
    return fields


def _field_ranges(fields: list[tuple[str, Any | None, str]]) -> list[tuple[int, int]]:
    ranges: list[tuple[int, int]] = []
    cursor = 0
    for index, (_kind, _owner, value) in enumerate(fields):
        ranges.append((cursor, cursor + len(value)))
        cursor += len(value)
        if index + 1 < len(fields):
            cursor += 1
    return ranges


def _safe_sheet_title(workbook: Any, sheet: Any, value: str) -> str:
    """Return a valid, unique Excel title after masking sensitive text."""
    title = re.sub(r"[\\/*?:\[\]]", "_", value).strip("'")[:31]
    if not title:
        title = "已脱敏工作表"
    existing = {item.title.casefold() for item in workbook.worksheets if item is not sheet}
    candidate = title
    counter = 2
    while candidate.casefold() in existing:
        suffix = f"_{counter}"
        candidate = title[: 31 - len(suffix)] + suffix
        counter += 1
    return candidate


def _safe_defined_name(value: str) -> str:
    """Normalize a masked value to Excel's defined-name character set."""
    name = re.sub(r"[^\w.\\]", "_", value, flags=re.UNICODE)[:255]
    if not name or not re.match(r"[A-Za-z_\\]", name):
        name = "_" + name
    return name[:255]


def _replace_sheet_reference(value: str, old_title: str, new_title: str) -> str:
    """Update formulas/defined names after a reviewed worksheet rename."""
    from openpyxl.utils.cell import quote_sheetname

    result = value.replace(f"{quote_sheetname(old_title)}!", f"{quote_sheetname(new_title)}!")
    result = result.replace(f"{old_title}!", f"{quote_sheetname(new_title)}!")

    # The title and its formula occurrence may both have been explicitly
    # reviewed.  In that case the formula temporarily contains masking
    # characters while the actual sheet title has normalized Excel's illegal
    # characters to underscores.  Reconcile those equivalent tokens after all
    # reviewed replacements have been applied.
    def replace_quoted(match: re.Match[str]) -> str:
        title = match.group(1).replace("''", "'")
        normalized = re.sub(r"[\\/*?:\[\]]", "_", title).strip("'")[:31]
        return f"{quote_sheetname(new_title)}!" if normalized == new_title else match.group(0)

    result = re.sub(r"'((?:''|[^'])*)'!", replace_quoted, result)

    def replace_unquoted(match: re.Match[str]) -> str:
        title = match.group(1)
        normalized = re.sub(r"[\\/*?:\[\]]", "_", title).strip("'")[:31]
        return f"{quote_sheetname(new_title)}!" if normalized == new_title else match.group(0)

    result = re.sub(r"(?<![\w.])([^\s=+*/^&<>,():']+)!", replace_unquoted, result)
    return result


def _formula_string_literal_mask(value: str, start: int, end: int) -> bool:
    """Return whether a formula span is inside an Excel double-quoted literal.

    Excel escapes a quote inside a literal by doubling it (``""``).  A small
    scanner is more reliable here than a regular expression because formulas
    can contain escaped quotes and nested function calls.  The offsets are
    always measured against the original formula string.
    """
    if not value.startswith("=") or start < 0 or end <= start or end > len(value):
        return False
    quoted = False
    index = 0
    while index < start:
        if value[index] == '"':
            if quoted and index + 1 < start and value[index + 1] == '"':
                index += 2
                continue
            quoted = not quoted
        index += 1
    if not quoted:
        return False
    # A reviewed span that crosses a literal boundary is not safe to replay as
    # a string replacement.  Treat it as an outside-literal token instead.
    while index < end:
        if value[index] == '"':
            if index + 1 < end and value[index + 1] == '"':
                index += 2
                continue
            return False
        index += 1
    return True


def _formula_safe_replacement(replacement: str, entity: dict[str, Any], value: str,
                              start: int, end: int) -> str:
    """Keep a replacement syntactically valid when it is outside a literal.

    Sensitive values may occur as a numeric operand or as part of a worksheet
    reference.  The normal text replacement intentionally contains masking
    characters, which are not valid in those positions.  Preserve valid Excel
    identifier/number tokens; otherwise use a neutral numeric constant.  The
    original value is never retained in the formula.
    """
    if _formula_string_literal_mask(value, start, end):
        # Keep the enclosing Excel literal valid when a custom policy includes
        # a quote or line break.
        return replacement.replace('"', '""').replace('\r', ' ').replace('\n', ' ')
    # A single-quoted worksheet name is another string-like Excel grammar
    # region.  Masking characters are legal there, and later sheet-reference
    # reconciliation will point it at the normalized title.
    quoted_sheet_start = value.rfind("'", 0, start + 1)
    quoted_sheet_end = value.find("'!", end)
    if quoted_sheet_start >= 0 and quoted_sheet_end >= end:
        return replacement.replace("'", "''").replace('\r', ' ').replace('\n', ' ')
    if re.fullmatch(r"(?:[A-Za-z_\\][A-Za-z0-9_.\\]*|[+-]?(?:\d+(?:\.\d*)?|\.\d+))", replacement):
        return replacement
    # A replacement containing a quoted string is valid in function arguments,
    # but not in a sheet/range token.  A neutral constant is valid in both
    # contexts and keeps the workbook loadable even for malformed source
    # formulas.
    return "0"


def _apply_formula_entities(value: str, entities: list[dict], policies: dict | None,
                            mapping: dict[tuple[str, str], str]) -> tuple[str, dict]:
    """Apply reviewed entities to a formula while preserving formula syntax."""
    # ``apply_entities`` performs the canonical offset, overlap and source
    # verification used by every other adapter.  It also populates the shared
    # replacement mapping, so identical values remain deterministic across a
    # workbook.  We then replay the replacements with context-aware tokens.
    apply_entities(value, entities, policies, mapping)
    selected = [entity for entity in entities if entity.get("selected", True)]
    ordered = sorted(selected, key=lambda item: (int(item["start"]), int(item["end"])), reverse=True)
    output = value
    for entity in ordered:
        start, end = int(entity["start"]), int(entity["end"])
        replacement = replacement_for(entity, policies, mapping)
        replacement = _formula_safe_replacement(replacement, entity, value, start, end)
        output = output[:start] + replacement + output[end:]
    return output, mapping


def _replace_xlsx(data: bytes, entities: list[dict], policies: dict | None, filename: str) -> bytes:
    from openpyxl import load_workbook

    keep_vba = filename.lower().endswith(".xlsm")
    workbook = load_workbook(
        io.BytesIO(data),
        read_only=False,
        data_only=False,
        keep_vba=keep_vba,
        keep_links=True,
    )
    fields = _xlsx_fields(workbook)
    ranges = _field_ranges(fields)
    mapping: dict[tuple[str, str], str] = {}
    sheet_renames: list[tuple[str, str]] = []
    defined_name_renames: list[tuple[str, Any, str]] = []
    for field_index, (kind, owner, value) in enumerate(fields):
        field_start, field_end = ranges[field_index]
        local_entities: list[dict] = []
        for entity in entities:
            target = str(entity.get("text", ""))
            if not target:
                continue
            if "start" in entity and "end" in entity:
                try:
                    start, end = int(entity["start"]), int(entity["end"])
                except (TypeError, ValueError):
                    continue
                if field_start <= start and end <= field_end and end > start:
                    local = dict(entity)
                    local["start"], local["end"] = start - field_start, end - field_start
                    local_entities.append(local)
            else:
                # Backwards-compatible direct use of this adapter.  The API
                # always sends offsets, so this fallback is deliberately opt-in
                # and does not replace all equal strings when review selected a
                # single occurrence.
                local_entities.extend(
                    {**entity, "start": match.start(), "end": match.end()}
                    for match in re.finditer(re.escape(target), value)
                )
        if not local_entities:
            continue
        if kind == "formula":
            masked, mapping = _apply_formula_entities(value, local_entities, policies, mapping)
        else:
            masked, mapping = apply_entities(value, local_entities, policies, mapping)
        if kind == "sheet" and owner is not None:
            old_title = str(owner.title)
            title_value = masked[1:-1] if masked.startswith("[") and masked.endswith("]") else masked
            new_title = _safe_sheet_title(workbook, owner, title_value)
            owner.title = new_title
            if old_title != new_title:
                sheet_renames.append((old_title, new_title))
        elif kind in {"value", "formula"} and owner is not None:
            owner.value = masked
        elif kind == "comment" and owner is not None and owner.comment is not None:
            owner.comment.text = masked
        elif kind == "url" and owner is not None and owner.hyperlink is not None:
            owner.hyperlink.target = masked
        elif kind == "header_footer" and owner is not None:
            header_footer, position = owner
            getattr(header_footer, position).text = masked
        elif kind == "defined_name" and owner is not None:
            defined_name_renames.append((str(owner.name), owner, _safe_defined_name(masked)))
        elif kind == "defined_value" and owner is not None:
            owner.attr_text = masked

    # A renamed sheet must remain addressable from formulas and defined-name
    # references.  Update those links even when the linked occurrence was not
    # independently selected in the review list.
    for old_title, new_title in sheet_renames:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        cell.value = _replace_sheet_reference(cell.value, old_title, new_title)
        for defined in workbook.defined_names.values():
            if isinstance(getattr(defined, "attr_text", None), str):
                defined.attr_text = _replace_sheet_reference(defined.attr_text, old_title, new_title)

    # ``DefinedNameDict`` is keyed by the original name.  Re-key selected
    # definitions after all fields have been processed so serialization and
    # subsequent lookups agree with the masked name.
    for old_name, defined, new_name in defined_name_renames:
        if workbook.defined_names.get(old_name) is defined:
            del workbook.defined_names[old_name]
        defined.name = new_name
        suffix = 2
        candidate = new_name
        while candidate in workbook.defined_names:
            candidate = f"{new_name[:250]}_{suffix}"
            suffix += 1
        defined.name = candidate
        workbook.defined_names.add(defined)
    _clear_xlsx_metadata(workbook)
    output = io.BytesIO()
    try:
        workbook.save(output)
    finally:
        # ``keep_vba`` stores a private ZipFile on the workbook.  Closing it
        # here prevents a delayed ``ZipFile.__del__`` warning on Windows and
        # releases the in-memory copy of the macro package promptly.
        vba_archive = getattr(workbook, "vba_archive", None)
        if vba_archive is not None:
            vba_archive.close()
    return output.getvalue()


def _clear_xlsx_metadata(workbook: Any) -> None:
    """Clear document properties and print metadata from an Excel output."""
    properties = getattr(workbook, "properties", None)
    if properties is not None:
        for attribute in (
            "creator", "lastModifiedBy", "title", "description", "subject",
            "identifier", "language", "keywords", "category", "contentStatus",
            "version", "revision",
        ):
            try:
                setattr(properties, attribute, None)
            except (AttributeError, TypeError, ValueError):
                pass
        # openpyxl's W3CDTF serializer requires created/modified to remain
        # datetimes.  Use a stable neutral value instead of retaining source
        # timestamps; ``lastPrinted`` is optional and can be removed.
        neutral_timestamp = datetime(1980, 1, 1)
        for attribute in ("created", "modified"):
            try:
                setattr(properties, attribute, neutral_timestamp)
            except (AttributeError, TypeError, ValueError):
                pass
        try:
            properties.lastPrinted = None
        except (AttributeError, TypeError, ValueError):
            pass
    for sheet in getattr(workbook, "worksheets", []):
        for row in sheet.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    cell.comment.author = ""
    # Defined names may carry a description/comment in newer openpyxl builds.
    for defined in getattr(workbook, "defined_names", {}).values():
        for attribute in ("comment", "description", "help", "statusBar"):
            if hasattr(defined, attribute):
                try:
                    setattr(defined, attribute, None)
                except (AttributeError, TypeError, ValueError):
                    pass


def _pdf_word_ranges(page: Any, page_text: str) -> list[tuple[int, int, Any]]:
    """Map character ranges in ``page.get_text()`` to word rectangles."""
    words = page.get_text("words", sort=True) or []
    result: list[tuple[int, int, Any]] = []
    cursor = 0
    for word in words:
        if len(word) < 5:
            continue
        token = str(word[4] or "")
        if not token:
            continue
        start = page_text.find(token, cursor)
        if start < 0:
            # Ligatures and unusual extraction order can make the sequential
            # search fail.  A second search still gives a useful rectangle;
            # duplicate tokens are disambiguated by the first pass whenever
            # possible.
            start = page_text.find(token)
        if start < 0:
            continue
        end = start + len(token)
        result.append((start, end, word))
        cursor = end
    return result


def _pdf_entity_span(entity: dict, page_number: int, page_text: str, page_start: int, page_end: int) -> tuple[int, int] | None:
    target = str(entity.get("text", ""))
    if not target:
        return None
    try:
        start = int(entity.get("start", -1))
        end = int(entity.get("end", -1))
    except (TypeError, ValueError):
        start = end = -1
    # Analyzer offsets are relative to the newline-joined document.  OCR and
    # hand-created entities commonly use page-local offsets, so accept either
    # representation when the text verifies exactly.
    if page_start <= start and end <= page_end and end > start:
        local_start, local_end = start - page_start, end - page_start
        if page_text[local_start:local_end] == target:
            return local_start, local_end
    if 0 <= start < end <= len(page_text) and page_text[start:end] == target:
        return start, end
    hint = max(0, start - page_start) if start >= page_start else max(0, start)
    found = page_text.find(target, hint)
    if found < 0:
        found = page_text.find(target)
    return (found, found + len(target)) if found >= 0 else None


def _pdf_bbox_rect(entity: dict, page: Any, fitz: Any) -> Any | None:
    bbox = entity.get("bbox")
    if not isinstance(bbox, dict):
        return None
    try:
        x = float(bbox.get("x", 0)); y = float(bbox.get("y", 0))
        width = float(bbox.get("width", 0)); height = float(bbox.get("height", 0))
    except (TypeError, ValueError):
        return None
    if not all(math.isfinite(value) for value in (x, y, width, height)):
        return None
    # Normalized boxes are the public protocol.  Accept absolute PDF points
    # as a convenience for OCR integrations that already use page coordinates.
    if 0 <= x <= 1 and 0 <= y <= 1 and 0 < width <= 1 and 0 < height <= 1:
        rect = page.rect
        return fitz.Rect(rect.x0 + x * rect.width, rect.y0 + y * rect.height,
                         rect.x0 + (x + width) * rect.width, rect.y0 + (y + height) * rect.height)
    if width <= 0 or height <= 0:
        return None
    return fitz.Rect(x, y, x + width, y + height)


_PDF_EMBEDDED_PATTERNS: tuple[tuple[str, re.Pattern[str]], ...] = (
    ("PHONE", re.compile(r"(?<!\d)(?:1[3-9]\d{9}|0\d{2,3}-?\d{7,8})(?!\d)")),
    ("ID_CARD", re.compile(r"(?<![0-9A-Za-z])\d{17}[\dXx](?![0-9A-Za-z])")),
    ("EMAIL", re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")),
    ("IP_ADDRESS", re.compile(r"(?<![\d.])(?:\d{1,3}\.){3}\d{1,3}(?![\d.])")),
)


def _mask_embedded_pdf_value(value: str, policies: dict | None,
                             mapping: dict[tuple[str, str], str]) -> str:
    """Remove common sensitive values from a non-page-text PDF field.

    URI annotations are not returned by ``page.get_text()``, so they cannot be
    selected using ordinary document offsets.  They still ship with the output
    artifact and therefore need an output-level privacy pass.  The patterns
    intentionally mirror the high-confidence local recognizers without
    importing ``app.main`` (which would create a circular dependency).
    """
    entities: list[dict[str, Any]] = []
    for entity_type, pattern in _PDF_EMBEDDED_PATTERNS:
        for match in pattern.finditer(value):
            if entity_type == "IP_ADDRESS":
                parts = match.group().split(".")
                if len(parts) != 4 or any(int(part) > 255 for part in parts):
                    continue
            entities.append({
                "type": entity_type,
                "text": match.group(),
                "start": match.start(),
                "end": match.end(),
                "selected": True,
            })
    # Match the analyzer's overlap rule so an e-mail containing another token
    # is replaced as one logical value.
    accepted: list[dict[str, Any]] = []
    for entity in sorted(entities, key=lambda item: (int(item["start"]), -int(item["end"]))):
        if not any(int(entity["start"]) < int(other["end"]) and int(other["start"]) < int(entity["end"])
                   for other in accepted):
            accepted.append(entity)
    return apply_entities(value, accepted, policies, mapping)[0] if accepted else value


def _redact_pdf_links(page: Any, fitz: Any, policies: dict | None,
                      mapping: dict[tuple[str, str], str]) -> None:
    """Sanitize URI annotations while preserving link rectangles."""
    for link in list(page.get_links() or []):
        if int(link.get("kind", 0) or 0) != int(fitz.LINK_URI):
            continue
        uri = str(link.get("uri", ""))
        masked_uri = _mask_embedded_pdf_value(uri, policies, mapping)
        if not uri or masked_uri == uri:
            continue
        # Updating the URI in place retains annotation flags and appearance.
        # Keep a delete/reinsert fallback for older PyMuPDF builds that reject
        # unusual replacement characters in an action dictionary.
        updated = dict(link)
        updated["uri"] = masked_uri
        try:
            page.update_link(updated)
        except (RuntimeError, ValueError):
            rectangle = link.get("from")
            page.delete_link(link)
            if rectangle is not None:
                page.insert_link({"kind": fitz.LINK_URI, "from": rectangle, "uri": masked_uri})


def _redact_pdf(
    data: bytes,
    entities: list[dict],
    policies: dict | None,
    boxes: list[dict] | None = None,
) -> bytes:
    import fitz

    document = fitz.open(stream=data, filetype="pdf")
    page_texts = [page.get_text() for page in document]
    page_ranges: list[tuple[int, int]] = []
    cursor = 0
    for index, value in enumerate(page_texts):
        page_ranges.append((cursor, cursor + len(value)))
        cursor += len(value)
        if index + 1 < len(page_texts):
            cursor += 1
    mapping: dict[tuple[str, str], str] = {}
    # OCR/manual boxes are first-class review regions. Older callers only
    # supplied entities, so retain that API while allowing the unified file
    # endpoint to redact a box even when it contains no recognized PII.
    all_entities: list[dict] = [*entities, *(boxes or [])]
    for page_index, page in enumerate(document):
        page_number = page_index + 1
        page_text = page_texts[page_index]
        page_start, page_end = page_ranges[page_index]
        word_ranges = _pdf_word_ranges(page, page_text)
        annotations: list[tuple[Any, str]] = []
        seen_regions: set[tuple[Any, ...]] = set()
        for entity in all_entities:
            requested_page = entity.get("page")
            if requested_page not in (None, "", 0):
                try:
                    if int(requested_page) != page_number:
                        continue
                except (TypeError, ValueError):
                    continue
            else:
                try:
                    entity_start = int(entity.get("start", -1))
                except (TypeError, ValueError):
                    entity_start = -1
                if not (page_start <= entity_start < page_end):
                    continue
            if not entity.get("selected", True):
                continue
            policy = policy_for(str(entity.get("type", "DEFAULT")), policies)
            if policy.get("text_action", "replace") == "keep":
                continue
            span = _pdf_entity_span(entity, page_number, page_text, page_start, page_end)
            rectangles: list[Any] = []
            if span is not None:
                start, end = span
                covered = [word for word_start, word_end, word in word_ranges if word_start < end and word_end > start]
                if covered:
                    rectangles = [fitz.Rect(word[0], word[1], word[2], word[3]) for word in covered]
            if not rectangles:
                bbox_rect = _pdf_bbox_rect(entity, page, fitz)
                if bbox_rect is not None:
                    rectangles = [bbox_rect]
            if not rectangles:
                continue
            # A single annotation over all words ensures that a value split by
            # spaces or line wrapping is removed as one logical entity.
            rect = rectangles[0]
            for other in rectangles[1:]:
                rect |= other
            # An OCR line and the sensitive entity found inside it may resolve
            # to the same rectangle.  De-duplicate by geometry (entities are
            # iterated first), so a generic OCR box cannot overwrite the
            # entity's replacement text.
            region_key = (
                page_number,
                round(float(rect.x0), 4), round(float(rect.y0), 4),
                round(float(rect.x1), 4), round(float(rect.y1), 4),
            )
            if region_key in seen_regions:
                continue
            seen_regions.add(region_key)
            replacement = replacement_for(entity, policies, mapping)
            annotations.append((rect, replacement))
        for rect, replacement in annotations:
            page.add_redact_annot(rect, text=replacement, fill=(0, 0, 0), text_color=(1, 1, 1))
        if annotations:
            page.apply_redactions()
        _redact_pdf_links(page, fitz, policies, mapping)
    output = io.BytesIO()
    # Remove document metadata and garbage-collect deleted text streams so a
    # downstream parser cannot recover the original content from the file.
    document.set_metadata({})
    try:
        document.set_xml_metadata("")
    except (AttributeError, RuntimeError, ValueError):
        pass
    document.save(output, garbage=4, deflate=True)
    document.close()
    return output.getvalue()


def _image_color(value: Any) -> tuple[int, int, int, int]:
    raw = str(value or "#000000").strip().lstrip("#")
    if len(raw) == 3:
        raw = "".join(char * 2 for char in raw)
    if len(raw) not in {6, 8} or not re.fullmatch(r"[0-9a-fA-F]+", raw):
        raw = "000000"
    if len(raw) == 6:
        raw += "ff"
    try:
        return tuple(int(raw[index:index + 2], 16) for index in range(0, 8, 2))  # type: ignore[return-value]
    except ValueError:
        return (0, 0, 0, 255)


def _image_box(entity: dict) -> tuple[float, float, float, float]:
    bbox = entity.get("bbox") if isinstance(entity, dict) else None
    bbox = bbox if isinstance(bbox, dict) else entity
    try:
        x = float(bbox.get("x", 0)); y = float(bbox.get("y", 0))
        width = float(bbox.get("width", 0)); height = float(bbox.get("height", 0))
    except (TypeError, ValueError) as exc:
        raise ValueError("图片框坐标无效") from exc
    values = (x, y, width, height)
    if not all(math.isfinite(value) for value in values):
        raise ValueError("图片框坐标无效")
    if x < 0 or y < 0 or width <= 0 or height <= 0 or x + width > 1 or y + height > 1:
        raise ValueError("图片框坐标必须位于 0 到 1 范围内")
    return x, y, width, height


def _draw_mask_text(image: Image.Image, box: tuple[int, int, int, int], value: str) -> None:
    left, top, right, _bottom = box
    draw = ImageDraw.Draw(image)
    draw.rectangle(box, fill=(255, 255, 255, 255))
    # A bundled font is not required.  The default bitmap font is clipped to
    # the box so a long replacement cannot spill over neighboring content.
    try:
        font = ImageFont.load_default()
    except (OSError, AttributeError):
        font = None
    text = str(value or "已脱敏")
    if font is not None:
        while text and draw.textbbox((0, 0), text, font=font)[2] > max(1, right - left - 4):
            text = text[:-1]
        draw.text((left + 2, top + 2), text, fill=(0, 0, 0, 255), font=font)


def _redact_image(
    data: bytes,
    entities: list[dict],
    filename: str,
    policies: dict | None = None,
    boxes: list[dict] | None = None,
) -> bytes:
    with Image.open(io.BytesIO(data)) as opened:
        # Apply EXIF orientation before calculating normalized boxes and never
        # copy metadata into the output artifact.
        image = ImageOps.exif_transpose(opened).convert("RGBA")
    reviewed_boxes = list(boxes or [])
    reviewed_ids = {
        str(box.get("id"))
        for box in reviewed_boxes
        if isinstance(box, dict) and box.get("id")
    }
    reviewed_regions = {
        tuple(round(value, 8) for value in _image_box(box))
        for box in reviewed_boxes
        if isinstance(box, dict)
    }
    # OCR entities carry the same geometry as their parent OCR line.  When the
    # reviewed box is present it is the user's authoritative region and image
    # policy; applying the entity first would turn "keep" or "blur" into the
    # entity type's default solid mask.  Retain standalone/manual entities.
    standalone_entities: list[dict] = []
    for entity in entities:
        linked_ids = {str(value) for value in entity.get("box_ids", []) if value}
        linked = bool(linked_ids & reviewed_ids)
        if not linked and isinstance(entity.get("bbox"), dict):
            linked = tuple(round(value, 8) for value in _image_box(entity)) in reviewed_regions
        if not linked:
            standalone_entities.append(entity)
    all_boxes = [*standalone_entities, *reviewed_boxes]
    seen_regions: set[tuple[float, float, float, float]] = set()
    for entity in all_boxes:
        # The review payload carries the user's explicit selection for both
        # OCR entities and manually drawn boxes.  A deselected region must be
        # left pixel-for-pixel untouched; metadata normalization alone should
        # not count as masking.
        if not entity.get("selected", True):
            continue
        x, y, width, height = _image_box(entity)
        region_key = tuple(round(value, 8) for value in (x, y, width, height))
        if region_key in seen_regions:
            continue
        seen_regions.add(region_key)
        left = max(0, min(image.width - 1, int(x * image.width)))
        top = max(0, min(image.height - 1, int(y * image.height)))
        right = min(image.width, max(left + 1, int((x + width) * image.width)))
        bottom = min(image.height, max(top + 1, int((y + height) * image.height)))
        if right <= left or bottom <= top:
            continue
        policy = policy_for(str(entity.get("type", "DEFAULT")), policies)
        action = str(policy.get("image_action", "solid"))
        if action == "keep":
            continue
        region = image.crop((left, top, right, bottom))
        if action == "blur":
            try:
                radius = max(1, min(80, int(policy.get("blur_radius", 14))))
            except (TypeError, ValueError):
                radius = 14
            image.alpha_composite(region.filter(ImageFilter.GaussianBlur(radius)), (left, top))
            if policy.get("show_replacement", True):
                replacement = replacement_for(entity, policies, {})
                _draw_mask_text(image, (left, top, right, bottom), replacement)
        elif action == "pixelate":
            try:
                factor = max(2, min(40, int(policy.get("pixel_size", 12) or 12)))
            except (TypeError, ValueError):
                factor = 12
            small = region.resize((max(1, region.width // factor), max(1, region.height // factor)), Image.Resampling.BILINEAR)
            image.alpha_composite(small.resize(region.size, Image.Resampling.NEAREST), (left, top))
        elif action == "text":
            _draw_mask_text(image, (left, top, right, bottom), str(policy.get("replacement") or "已脱敏"))
        else:
            image.paste(_image_color(policy.get("color", "#000000")), (left, top, right, bottom))
    output = io.BytesIO()
    ext = extension(filename)
    save_format = {"jpg": "JPEG", "jpeg": "JPEG", "bmp": "BMP", "tif": "TIFF", "tiff": "TIFF", "png": "PNG"}.get(ext, "PNG")
    if save_format in {"JPEG", "BMP", "TIFF"}:
        image = image.convert("RGB")
    # Explicitly omit EXIF/ICC/comment metadata.  This is important for a
    # desensitization artifact even when no boxes were selected.
    save_kwargs: dict[str, Any] = {}
    if save_format == "JPEG":
        save_kwargs.update(quality=95, optimize=True, progressive=False)
    image.save(output, format=save_format, **save_kwargs)
    return output.getvalue()


def mask_file(data: bytes, filename: str, text: str, entities: list[dict], policies: dict | None = None, boxes: list[dict] | None = None) -> bytes:
    ext = extension(filename)
    if ext in {"txt", "md"}:
        return apply_entities(text, entities, policies)[0].encode("utf-8")
    if ext == "docx":
        return _replace_docx_text(data, entities, policies)
    if ext in {"xlsx", "xlsm"}:
        return _replace_xlsx(data, entities, policies, filename)
    if ext == "pdf":
        return _redact_pdf(data, entities, policies, boxes)
    if ext in {"png", "jpg", "jpeg", "bmp", "tif", "tiff"}:
        return _redact_image(data, entities, filename, policies, boxes)
    raise ValueError(f"不支持的文件格式: .{ext}")


def content_manifest(filename: str, text: str, entities: list[dict]) -> dict:
    return {"filename": filename, "text_length": len(text), "entity_count": len(entities), "entities": entities}

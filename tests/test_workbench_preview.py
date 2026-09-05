"""Workbench preview and browser-shell acceptance contracts.

The HTTP tests use real file containers and masking adapters. Model inference
is replaced with deterministic local detections so the suite stays fast and
does not depend on downloaded weights.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageChops, ImageDraw

import app.main as main


PHONE = "13800138000"


def _phone_entities(text: str) -> list[dict]:
    start = text.find(PHONE)
    if start < 0:
        return []
    return [
        {
            "id": f"phone-{start}",
            "type": "PHONE",
            "text": PHONE,
            "start": start,
            "end": start + len(PHONE),
            "score": 0.99,
            "selected": True,
            "source": "fixture",
            "box_ids": [],
        }
    ]


def _write_fixture(path: Path) -> None:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        path.write_text(f"联系人电话 {PHONE}", encoding="utf-8")
        return
    if suffix == ".docx":
        document = Document()
        document.add_heading("客户联系表", level=1)
        document.add_paragraph(f"联系人电话 {PHONE}")
        document.save(path)
        return
    if suffix == ".xlsx":
        workbook = Workbook()
        workbook.active.title = "客户联系表"
        workbook.active["A1"] = "联系人电话"
        workbook.active["B1"] = PHONE
        workbook.save(path)
        return
    if suffix == ".pdf":
        document = fitz.open()
        document.new_page(width=595, height=842).insert_text(
            (72, 90), f"Contact phone: {PHONE}", fontsize=12
        )
        document.save(path)
        document.close()
        return
    raise AssertionError(f"unsupported fixture: {path}")


def _assert_phone_removed(filename: str, content: bytes) -> None:
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        assert PHONE not in content.decode("utf-8")
    elif suffix == ".docx":
        document = Document(io.BytesIO(content))
        assert PHONE not in "\n".join(paragraph.text for paragraph in document.paragraphs)
    elif suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        try:
            values = [cell.value for row in workbook.active.iter_rows() for cell in row]
            assert all(PHONE not in str(value or "") for value in values)
        finally:
            workbook.close()
    elif suffix == ".pdf":
        document = fitz.open(stream=content, filetype="pdf")
        try:
            assert PHONE not in "\n".join(page.get_text() for page in document)
        finally:
            document.close()
    else:
        raise AssertionError(f"unsupported output: {filename}")


def _assert_preview_view_contains_phone(view: dict, suffix: str) -> None:
    if suffix == ".txt":
        assert view["type"] == "text"
        assert PHONE in view["text"]
    elif suffix == ".docx":
        assert view["type"] == "document"
        assert any(PHONE in block.get("text", "") for block in view["blocks"])
    elif suffix == ".xlsx":
        assert view["type"] == "spreadsheet"
        assert any(
            PHONE in str(cell.get("value", ""))
            for sheet in view["sheets"]
            for cell in sheet["cells"]
        )
    elif suffix == ".pdf":
        assert view["type"] == "pdf"
        assert view["page_count"] == 1
        assert "{page}" in view["page_url_template"]
    else:
        raise AssertionError(f"unsupported preview: {suffix}")


@pytest.mark.parametrize(
    ("suffix", "media_type"),
    [
        (".txt", "text/plain"),
        (".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        (".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        (".pdf", "application/pdf"),
    ],
)
def test_real_document_preview_can_be_reopened_and_is_non_destructive(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    suffix: str,
    media_type: str,
) -> None:
    monkeypatch.setattr(main, "analyze", _phone_entities)
    source = tmp_path / f"客户资料{suffix}"
    _write_fixture(source)
    original = source.read_bytes()
    client = TestClient(main.app)

    with source.open("rb") as stream:
        analyzed_response = client.post(
            "/api/files/analyze",
            files={"file": (source.name, stream, media_type)},
        )
    assert analyzed_response.status_code == 200, analyzed_response.text
    analyzed = analyzed_response.json()
    task_id = analyzed["analysis_id"]

    reopened_response = client.get(f"/api/tasks/{task_id}/analysis")
    assert reopened_response.status_code == 200, reopened_response.text
    reopened = reopened_response.json()
    assert reopened["filename"] == source.name
    assert reopened["source_preview"] == f"/api/tasks/{task_id}/source-preview"
    assert any(item["type"] == "PHONE" for item in reopened["entities"])
    _assert_preview_view_contains_phone(reopened["source_view"], suffix)

    source_preview = client.get(reopened["source_preview"])
    assert source_preview.status_code == 200
    assert source_preview.headers["content-type"].startswith(media_type)
    assert source_preview.content == original

    preview_response = client.post(
        "/api/files/preview",
        json={
            "analysis_id": task_id,
            "filename": reopened["filename"],
            "text": reopened.get("text", ""),
            "entities": reopened["entities"],
            "boxes": reopened.get("boxes", []),
            "policies": {
                "PHONE": {
                    "text_action": "replace",
                    "replacement": "MASKED",
                }
            },
        },
    )
    assert preview_response.status_code == 200, preview_response.text
    preview = preview_response.json()
    assert preview["preview_url"].startswith(f"/api/tasks/{task_id}/previews/")
    assert preview["source_view"]["type"] == reopened["source_view"]["type"]
    assert preview["masked_view"]["type"] == reopened["source_view"]["type"]
    preview_file = client.get(preview["preview_url"])
    assert preview_file.status_code == 200
    assert preview_file.headers["content-type"].startswith(media_type)
    _assert_phone_removed(source.name, preview_file.content)

    task = next(item for item in client.get("/api/tasks").json()["tasks"] if item["id"] == task_id)
    assert task["status"] == "awaiting_review"
    assert not task["artifact"]
    assert task["has_preview"] is True
    assert client.get(f"/api/tasks/{task_id}/source-preview").content == original
    assert client.get(f"/api/tasks/{task_id}/analysis").status_code == 200
    if suffix == ".pdf":
        source_page = client.get(reopened["source_view"]["page_url_template"].format(page=1))
        masked_page = client.get(preview["masked_view"]["page_url_template"].format(page=1))
        assert source_page.status_code == 200
        assert masked_page.status_code == 200
        assert source_page.headers["content-type"].startswith("image/png")
        assert masked_page.headers["content-type"].startswith("image/png")
        assert source_page.content != masked_page.content


def test_real_image_preview_preserves_source_and_changes_selected_region(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    image = Image.new("RGB", (240, 120), "white")
    ImageDraw.Draw(image).rectangle((48, 24, 192, 84), fill=(20, 90, 180))
    source = tmp_path / "客户截图.png"
    image.save(source)
    original = source.read_bytes()

    box = {
        "id": "ocr-line-1",
        "type": "OCR",
        "text": PHONE,
        "page": 1,
        "x": 0.2,
        "y": 0.2,
        "width": 0.6,
        "height": 0.5,
        "bbox": {"x": 0.2, "y": 0.2, "width": 0.6, "height": 0.5},
        "entity_ids": ["image-phone"],
        "selected": True,
        "source": "ocr",
        "score": 0.99,
    }
    entity = {
        "id": "image-phone",
        "type": "PHONE",
        "text": PHONE,
        "start": 0,
        "end": len(PHONE),
        "score": 0.99,
        "selected": True,
        "source": "ocr",
        "page": 1,
        "bbox": dict(box["bbox"]),
        "box_ids": [box["id"]],
    }
    monkeypatch.setattr(
        main,
        "_analyze_image_content",
        lambda data, folder: (240, 120, [dict(box)], [dict(entity)], []),
    )
    client = TestClient(main.app)
    with source.open("rb") as stream:
        analyzed_response = client.post(
            "/api/files/analyze",
            files={"file": (source.name, stream, "image/png")},
        )
    assert analyzed_response.status_code == 200, analyzed_response.text
    analyzed = analyzed_response.json()
    task_id = analyzed["analysis_id"]
    assert analyzed["source_view"] == {
        "type": "image",
        "url": f"/api/tasks/{task_id}/source-image",
        "width": 240,
        "height": 120,
    }

    preview_response = client.post(
        "/api/files/preview",
        json={
            "analysis_id": task_id,
            "filename": source.name,
            "text": analyzed["text"],
            "entities": analyzed["entities"],
            "boxes": analyzed["boxes"],
            "image_override": True,
            "policies": {
                "PHONE": {
                    "text_action": "mask",
                    "image_action": "solid",
                    "color": "#ff0000",
                    "show_replacement": False,
                }
            },
        },
    )
    assert preview_response.status_code == 200, preview_response.text
    preview_file = client.get(preview_response.json()["preview_url"])
    assert preview_file.status_code == 200
    rendered = Image.open(io.BytesIO(preview_file.content)).convert("RGB")
    assert ImageChops.difference(image, rendered).getbbox() is not None
    assert rendered.getpixel((100, 50)) == (255, 0, 0)

    preview = preview_response.json()
    assert preview["masked_view"]["type"] == "image"
    assert preview["masked_view"]["url"].endswith("/image")
    browser_image = client.get(preview["masked_view"]["url"])
    assert browser_image.status_code == 200
    assert browser_image.headers["content-type"].startswith("image/png")

    source_preview = client.get(f"/api/tasks/{task_id}/source-preview")
    assert source_preview.status_code == 200
    assert source_preview.content == original
    assert client.get(f"/api/tasks/{task_id}/analysis").json()["boxes"][0]["selected"] is True


def test_frontend_exposes_complete_local_workbench_contract() -> None:
    html = Path("app/static/index.html").read_text(encoding="utf-8")

    assert 'lang="zh-CN"' in html
    for label in (
        "文件中心",
        "脱敏工作台",
        "原文件",
        "脱敏后预览",
        "识别规则",
        "脱敏设置",
        "任务历史",
        "本地模型",
        "个人姓名",
        "电话号码",
        "执行脱敏",
    ):
        assert label in html

    for endpoint in (
        "/api/text/analyze",
        "/api/text/mask",
        "/api/files/analyze",
        "/api/files/preview",
        "/api/files/mask",
        "/api/tasks/",
        "/analysis",
        "/source-preview",
        "/api/custom-sensitive",
        "/api/policies",
        "/api/models/initialize",
    ):
        assert endpoint in html

    assert "showDirectoryPicker" in html
    assert "@media" in html
    assert "我已复核全部实体和页面区域" not in html
    assert "我已复核上方实体" not in html
    assert "<span hidden>" not in html
    assert not re.search(r">\s*(awaiting_review|completed|PHONE_NUMBER)\s*<", html)

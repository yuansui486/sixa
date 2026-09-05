"""Final-version workflow contracts using local fixtures only.

These tests intentionally replace model/OCR adapters with deterministic local
functions.  They exercise the HTTP contracts the UI relies on without
downloading or executing large models.
"""

from __future__ import annotations

import io
from pathlib import Path

import fitz
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image

import app.main as main


def _client() -> TestClient:
    return TestClient(main.app)


def _fake_entities(text: str) -> list[dict]:
    entities: list[dict] = []
    marker = "13800138000"
    start = text.find(marker)
    if start >= 0:
        entities.append(
            {
                "id": "fixture-phone",
                "type": "PHONE",
                "text": marker,
                "start": start,
                "end": start + len(marker),
                "score": 0.99,
                "selected": True,
            }
        )
    return entities


def test_entities_catalog_is_localized_and_analysis_can_be_reopened(monkeypatch) -> None:
    monkeypatch.setattr(main, "analyze", _fake_entities)
    client = _client()
    catalog = client.get("/api/entities")
    assert catalog.status_code == 200, catalog.text
    entries = catalog.json().get("entities", catalog.json().get("entity_catalog", []))
    assert entries
    assert any("电话" in str(item.get("name", item.get("label", ""))) for item in entries)

    analyzed = client.post(
        "/api/files/analyze",
        files={"file": ("resume.txt", "联系电话 13800138000".encode(), "text/plain")},
    )
    assert analyzed.status_code == 200, analyzed.text
    task_id = analyzed.json()["analysis_id"]
    reopened = client.get(f"/api/tasks/{task_id}/analysis")
    assert reopened.status_code == 200, reopened.text
    payload = reopened.json()
    assert payload["text"] == "联系电话 13800138000"
    assert payload["entities"][0]["type"] == "PHONE"


def test_preview_is_non_destructive_and_mask_produces_download(monkeypatch) -> None:
    monkeypatch.setattr(main, "analyze", _fake_entities)
    client = _client()
    analyzed = client.post(
        "/api/files/analyze",
        files={"file": ("preview.txt", "联系电话 13800138000".encode(), "text/plain")},
    ).json()
    task_id = analyzed["analysis_id"]

    preview = client.post(
        "/api/files/preview",
        json={
            "analysis_id": task_id,
            "filename": "preview.txt",
            "text": analyzed["text"],
            "entities": analyzed["entities"],
            "reviewed": True,
        },
    )
    assert preview.status_code == 200, preview.text
    assert preview.json().get("preview_url") or preview.json().get("masked_text")
    status = next(item for item in client.get("/api/tasks").json()["tasks"] if item["id"] == task_id)
    assert status["status"] == "awaiting_review"

    masked = client.post(
        "/api/files/mask",
        json={
            "analysis_id": task_id,
            "filename": "preview.txt",
            "text": analyzed["text"],
            "entities": analyzed["entities"],
            "reviewed": True,
        },
    )
    assert masked.status_code == 200, masked.text
    artifact = client.get(masked.json()["artifact"])
    assert artifact.status_code == 200
    assert b"13800138000" not in artifact.content


def test_custom_sensitive_literal_and_category_lifecycle() -> None:
    client = _client()
    literal = client.post(
        "/api/custom-sensitive",
        json={
            "mode": "literal",
            "value": "客户代号-ABC",
            "name": "客户代号",
            "entity_type": "CUSTOM",
            "replacement": "[客户代号]",
            "enabled": True,
        },
    )
    assert literal.status_code == 200, literal.text
    literal_id = literal.json()["id"]

    category = client.post(
        "/api/custom-sensitive",
        json={
            "mode": "category",
            "value": "PERSON",
            "name": "个人姓名",
            "entity_type": "PERSON",
            "replacement": "[姓名]",
            "enabled": True,
        },
    )
    assert category.status_code == 200, category.text
    items = client.get("/api/custom-sensitive").json()["items"]
    assert any(item["id"] == literal_id and item["mode"] == "literal" for item in items)
    assert any(item["mode"] == "category" and item["entity_type"] == "PERSON" for item in items)

    updated = client.patch(
        f"/api/custom-sensitive/{literal_id}",
        json={
            "mode": "literal",
            "value": "客户代号-XYZ",
            "name": "客户代号",
            "entity_type": "CUSTOM",
            "replacement": "[客户]",
            "enabled": False,
        },
    )
    assert updated.status_code == 200, updated.text
    deleted = client.delete(f"/api/custom-sensitive/{literal_id}")
    assert deleted.status_code == 200, deleted.text


def test_real_file_fixtures_analyze_without_model_execution(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setattr(main, "analyze", _fake_entities)
    monkeypatch.setattr(
        main,
        "_analyze_image_content",
        lambda data, folder: (100, 50, [{"text": "13800138000", "x": 0.1, "y": 0.1, "width": 0.5, "height": 0.3}], _fake_entities("13800138000"), []),
    )

    docx = Document()
    docx.add_paragraph("电话 13800138000")
    docx_buf = io.BytesIO()
    docx.save(docx_buf)
    xlsx = Workbook()
    xlsx.active["A1"] = "电话 13800138000"
    xlsx_buf = io.BytesIO()
    xlsx.save(xlsx_buf)
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "电话 13800138000")
    pdf_buf = io.BytesIO(pdf.tobytes())
    pdf.close()
    image_buf = io.BytesIO()
    Image.new("RGB", (100, 50), "white").save(image_buf, format="PNG")

    client = _client()
    fixtures = [
        ("fixture.docx", docx_buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        ("fixture.xlsx", xlsx_buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("fixture.pdf", pdf_buf.getvalue(), "application/pdf"),
        ("fixture.png", image_buf.getvalue(), "image/png"),
    ]
    for filename, content, mime in fixtures:
        response = client.post("/api/files/analyze", files={"file": (filename, content, mime)})
        assert response.status_code == 200, f"{filename}: {response.text}"
        body = response.json()
        assert body["analysis_id"]
        assert body["filename"] == filename


def test_preview_contract_renders_text_office_pdf_and_image(monkeypatch) -> None:
    monkeypatch.setattr(main, "analyze", _fake_entities)

    image_box = {
        "id": "fixture-box",
        "type": "OCR",
        "text": "13800138000",
        "x": 0.1,
        "y": 0.2,
        "width": 0.7,
        "height": 0.4,
        "page": 1,
        "source": "ocr",
        "selected": True,
        "entity_ids": ["fixture-phone"],
    }
    image_entity = {
        **_fake_entities("13800138000")[0],
        "bbox": {key: image_box[key] for key in ("x", "y", "width", "height")},
        "page": 1,
        "box_ids": ["fixture-box"],
    }
    monkeypatch.setattr(
        main,
        "_analyze_image_content",
        lambda _data, _folder: (240, 120, [image_box], [image_entity], []),
    )

    document = Document()
    document.add_heading("客户资料", level=1)
    document.add_paragraph("联系电话 13800138000")
    document.add_table(rows=1, cols=2).rows[0].cells[0].text = "表格内容"
    document.sections[0].header.paragraphs[0].text = "内部文件"
    docx_buffer = io.BytesIO()
    document.save(docx_buffer)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "客户"
    worksheet["A1"] = "联系电话"
    worksheet["B1"] = "13800138000"
    worksheet["C1"] = "=LEN(B1)"
    xlsx_buffer = io.BytesIO()
    workbook.save(xlsx_buffer)

    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Phone 13800138000")
    pdf_bytes = pdf.tobytes()
    pdf.close()

    image_buffer = io.BytesIO()
    Image.new("RGB", (240, 120), "white").save(image_buffer, format="PNG")

    fixtures = [
        ("note.txt", "联系电话 13800138000".encode(), "text/plain", "text"),
        (
            "resume.docx",
            docx_buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "document",
        ),
        (
            "clients.xlsx",
            xlsx_buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "spreadsheet",
        ),
        ("resume.pdf", pdf_bytes, "application/pdf", "pdf"),
        ("scan.png", image_buffer.getvalue(), "image/png", "image"),
    ]

    client = _client()
    for filename, content, mime, expected_type in fixtures:
        analyzed = client.post("/api/files/analyze", files={"file": (filename, content, mime)})
        assert analyzed.status_code == 200, f"{filename}: {analyzed.text}"
        analysis = analyzed.json()
        assert analysis["source_view"]["type"] == expected_type
        assert analysis["source_preview"]

        preview = client.post(
            "/api/files/preview",
            json={
                "analysis_id": analysis["analysis_id"],
                # Omitting filename is a supported way to use the immutable
                # server-side analysis filename.
                "entities": analysis.get("entities", []),
                "boxes": analysis.get("boxes", []),
                "image_override": filename.endswith(".png"),
            },
        )
        assert preview.status_code == 200, f"{filename}: {preview.text}"
        payload = preview.json()
        assert payload["source_view"]["type"] == expected_type
        assert payload["masked_view"]["type"] == expected_type

        reopened = client.get(f"/api/tasks/{analysis['analysis_id']}/analysis")
        assert reopened.status_code == 200, reopened.text
        assert reopened.json()["source_view"]["type"] == expected_type

        if expected_type == "image":
            assert client.get(analysis["source_view"]["url"]).headers["content-type"] == "image/png"
            assert client.get(payload["masked_view"]["url"]).headers["content-type"] == "image/png"
        elif expected_type == "pdf":
            source_page = analysis["source_view"]["page_url_template"].replace("{page}", "1")
            masked_page = payload["masked_view"]["page_url_template"].replace("{page}", "1")
            assert client.get(source_page).headers["content-type"] == "image/png"
            assert client.get(masked_page).headers["content-type"] == "image/png"

    doc_view = client.get(
        f"/api/tasks/{next(item['id'] for item in client.get('/api/tasks').json()['tasks'] if item['filename'] == 'resume.docx')}/analysis"
    ).json()["source_view"]
    assert any(block["type"] == "table" for block in doc_view["blocks"])
    assert any(block["location"] == "header" for block in doc_view["blocks"])

    sheet_view = client.get(
        f"/api/tasks/{next(item['id'] for item in client.get('/api/tasks').json()['tasks'] if item['filename'] == 'clients.xlsx')}/analysis"
    ).json()["source_view"]
    cells = sheet_view["sheets"][0]["cells"]
    assert any(cell["address"] == "C1" and cell["kind"] == "formula" for cell in cells)


def test_invalid_preview_page_and_tampered_preview_are_rejected(monkeypatch) -> None:
    monkeypatch.setattr(main, "analyze", _fake_entities)
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "Phone 13800138000")
    content = pdf.tobytes()
    pdf.close()
    client = _client()
    analysis = client.post(
        "/api/files/analyze",
        files={"file": ("resume.pdf", content, "application/pdf")},
    ).json()
    assert client.get(
        analysis["source_view"]["page_url_template"].replace("{page}", "2")
    ).status_code == 404

    preview = client.post(
        "/api/files/preview",
        json={
            "analysis_id": analysis["analysis_id"],
            "entities": analysis["entities"],
            "boxes": analysis.get("boxes", []),
        },
    ).json()
    preview_path = main.TASKS / analysis["analysis_id"] / ".previews" / (
        f"{preview['preview_id']}.pdf"
    )
    preview_path.write_bytes(preview_path.read_bytes() + b"tampered")
    page_url = preview["masked_view"]["page_url_template"].replace("{page}", "1")
    assert client.get(page_url).status_code == 409

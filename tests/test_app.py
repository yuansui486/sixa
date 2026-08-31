from io import BytesIO

from fastapi.testclient import TestClient
from PIL import Image, ImageChops, ImageDraw

from app.main import TASKS, app, cleanup, conn, luhn, valid_id, valid_ip

client = TestClient(app)

def test_validators():
    assert valid_id("11010519491231002X")
    assert not valid_id("110105194912310021")
    assert luhn("4532015112830366")
    assert not luhn("4532015112830367")
    assert valid_ip("192.168.1.1")
    assert not valid_ip("999.1.1.1")
    assert not valid_ip("1.2.3")
    assert not valid_ip("1.2.3.4.5")
    assert not valid_ip("abc.1.1.1")


def test_chinese_labels_do_not_hide_id_card_boundaries():
    text = "客户身份证11010519491231002X，信息已核验"
    entities = client.post("/api/text/analyze", json={"text": text}).json()["entities"]
    assert any(
        item["type"] == "ID_CARD" and item["text"] == "11010519491231002X"
        for item in entities
    )

def test_text_roundtrip_and_selection():
    text = "张三电话13800138000，邮箱 a@test.com"
    entities = client.post("/api/text/analyze", json={"text": text}).json()["entities"]
    entities[0]["selected"] = False
    result = client.post("/api/text/mask", json={"text": text, "entities": entities}).json()
    assert "13800138000" in result["masked_text"]
    assert "a@test.com" not in result["masked_text"]
    task = next(item for item in client.get("/api/tasks").json()["tasks"] if item["id"] == result["task_id"])
    assert task["status"] == "completed"
    restored = client.post("/api/text/unmask", json={"task_id": result["task_id"]}).json()
    assert restored["text"] == text

def test_invalid_entity_offsets():
    response = client.post("/api/text/mask", json={"text": "abc", "entities": [{"type": "X", "text": "bad", "start": 0, "end": 3}]})
    assert response.status_code == 422

def test_custom_rule_crud():
    created = client.post("/api/rules", json={"name": "internal", "kind": "word", "pattern": "ProjectX", "entity_type": "PROJECT"})
    assert created.status_code == 200
    entities = client.post("/api/text/analyze", json={"text": "ProjectX"}).json()["entities"]
    assert entities[0]["type"] == "PROJECT"
    assert client.delete(f"/api/rules/{created.json()['id']}").status_code == 200

def test_image_mask():
    image = Image.new("RGB", (100, 100), "white")
    ImageDraw.Draw(image).rectangle((10, 10, 60, 60), fill="black")
    ImageDraw.Draw(image).line((35, 10, 35, 60), fill="white", width=3)
    source = BytesIO(); image.save(source, format="PNG"); source.seek(0)
    analysis = client.post("/api/image/analyze", files={"file": ("a.png", source.getvalue(), "image/png")}).json()
    masked = client.post("/api/image/mask", json={"analysis_id": analysis["analysis_id"], "boxes": [{"x": 0.1, "y": 0.1, "width": 0.5, "height": 0.5}]}).json()
    assert masked["artifact"].endswith("masked.png")
    artifact = client.get(masked["artifact"])
    assert artifact.status_code == 200
    masked_image = Image.open(BytesIO(artifact.content)).convert("RGB")
    assert ImageChops.difference(image, masked_image).getbbox() is not None
    tasks = client.get("/api/tasks").json()["tasks"]
    assert any(t["id"] == analysis["analysis_id"] and t["kind"] == "image" for t in tasks)

def test_real_ocr_smoke():
    image = Image.new("RGB", (500, 120), "white")
    from PIL import ImageDraw
    ImageDraw.Draw(image).text((20, 40), "13800138000", fill="black")
    source = BytesIO(); image.save(source, format="PNG")
    response = client.post("/api/image/analyze", files={"file": ("ocr.png", source.getvalue(), "image/png")})
    assert response.status_code == 200
    payload = response.json()
    if payload["ocr_available"]:
        assert any("13800138000" in box["text"] for box in payload["boxes"])

def test_real_raner_smoke():
    from app.main import ner_service
    ner_service.load(download=False, device="cpu")
    if ner_service.available:
        entities = ner_service.analyze("张三在北京大学工作。")
        assert any(e["type"] in {"PERSON", "ORGANIZATION", "LOCATION", "GPE"} for e in entities)

def test_document_analyze_txt():
    response = client.post("/api/document/analyze", files={"file": ("a.txt", "电话13800138000".encode(), "text/plain")})
    assert response.status_code == 200
    payload = response.json()
    assert payload["entities"][0]["type"] == "PHONE"
    assert any(t["id"] == payload["analysis_id"] and t["kind"] == "document" for t in client.get("/api/tasks").json()["tasks"])
    masked = client.post("/api/document/mask", json={"analysis_id": payload["analysis_id"], "text": payload["text"], "entities": payload["entities"]}).json()
    assert masked["task_id"] == payload["analysis_id"]
    assert "13800138000" not in masked["masked_text"]
    assert client.get(masked["artifact"]).status_code == 200

def test_cleanup_removes_old_non_text_task():
    response = client.post("/api/document/analyze", files={"file": ("old.txt", b"old", "text/plain")})
    task_id = response.json()["analysis_id"]
    c = conn(); c.execute("UPDATE tasks SET created=?,updated=? WHERE id=?", (0, 0, task_id)); c.commit(); c.close()
    assert (TASKS / task_id).is_dir()
    cleanup()
    assert not (TASKS / task_id).exists()
    assert not any(t["id"] == task_id for t in client.get("/api/tasks").json()["tasks"])


def test_task_history_paginates_and_filters():
    created = []
    for index in range(3):
        response = client.post(
            "/api/text/mask",
            json={"text": f"paging-{index}", "entities": []},
        )
        assert response.status_code == 200
        task_id = response.json()["task_id"]
        c = conn()
        c.execute("UPDATE tasks SET kind=? WHERE id=?", ("paging_fixture", task_id))
        c.commit(); c.close()
        created.append(task_id)
    try:
        first = client.get("/api/tasks?kind=paging_fixture&limit=2&offset=0").json()
        second = client.get("/api/tasks?kind=paging_fixture&limit=2&offset=2").json()
        assert first["total"] == 3
        assert len(first["tasks"]) == 2
        assert len(second["tasks"]) == 1
        assert {item["id"] for item in first["tasks"] + second["tasks"]} == set(created)
    finally:
        for task_id in created:
            client.delete(f"/api/tasks/{task_id}")

def test_unified_docx_xlsx_pdf_image_and_batch():
    import fitz
    from docx import Document
    from openpyxl import Workbook, load_workbook

    doc = Document(); doc.add_paragraph("联系电话 13800138000")
    doc_bytes = BytesIO(); doc.save(doc_bytes)
    analyzed = client.post("/api/files/analyze", files={"file": ("sample.docx", doc_bytes.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}).json()
    masked = client.post("/api/files/mask", json={"analysis_id": analyzed["analysis_id"], "filename": "sample.docx", "text": analyzed["text"], "entities": analyzed["entities"], "reviewed": True}).json()
    assert client.get(masked["artifact"]).status_code == 200

    wb = Workbook(); ws = wb.active; ws["A1"] = "13800138000"; ws["B1"] = '="13800138000"'; xlsx = BytesIO(); wb.save(xlsx)
    analyzed = client.post("/api/files/analyze", files={"file": ("sample.xlsx", xlsx.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}).json()
    masked = client.post("/api/files/mask", json={"analysis_id": analyzed["analysis_id"], "filename": "sample.xlsx", "text": analyzed["text"], "entities": analyzed["entities"], "reviewed": True}).json()
    workbook = load_workbook(BytesIO(client.get(masked["artifact"]).content), data_only=False)
    assert "13800138000" not in str(workbook.active["A1"].value)

    pdf = fitz.open(); page = pdf.new_page(); page.insert_text((72, 72), "电话 13800138000"); pdf_bytes = pdf.tobytes()
    analyzed = client.post("/api/files/analyze", files={"file": ("sample.pdf", pdf_bytes, "application/pdf")}).json()
    masked = client.post("/api/files/mask", json={"analysis_id": analyzed["analysis_id"], "filename": "sample.pdf", "text": analyzed["text"], "entities": analyzed["entities"], "reviewed": True}).json()
    assert fitz.open(stream=client.get(masked["artifact"]).content, filetype="pdf").page_count == 1

    image = Image.new("RGB", (100, 100), "white"); source = BytesIO(); image.save(source, "PNG")
    analyzed = client.post("/api/files/analyze", files={"file": ("sample.png", source.getvalue(), "image/png")}).json()
    masked = client.post("/api/files/mask", json={"analysis_id": analyzed["analysis_id"], "filename": "sample.png", "text": "13800138000", "entities": [{"type": "PHONE", "text": "13800138000", "start": 0, "end": 11, "bbox": {"x": 0.1, "y": 0.1, "width": 0.5, "height": 0.5}, "selected": True}], "reviewed": True}).json()
    assert client.get(masked["artifact"]).status_code == 200

    batch = client.post("/api/batches", files=[("files", ("one.txt", b"13800138000", "text/plain")), ("files", ("two.txt", b"hello", "text/plain"))], data={"options": '{"reviewed": true}'}).json()
    import time
    for _ in range(100):
        status = client.get(f"/api/batches/{batch['batch_id']}").json()
        if status.get("status") == "awaiting_review":
            break
        time.sleep(0.05)
    assert status["status"] == "awaiting_review"
    reviewed_items = status.get("items", [])
    execute = client.post(f"/api/batches/{batch['batch_id']}/mask", json={"reviewed": True, "items": reviewed_items})
    assert execute.status_code == 200
    for _ in range(100):
        status = client.get(f"/api/batches/{batch['batch_id']}").json()
        if status.get("archive"): break
        time.sleep(0.05)
    assert status["status"] in {"completed", "partial"}
    assert client.get(f"/api/batches/{batch['batch_id']}/download").status_code == 200

def test_batch_review_gate_and_execute():
    batch = client.post("/api/batches", files=[("files", ("review.txt", "电话13800138000".encode(), "text/plain"))], data={"options": '{"reviewed": false}'}).json()
    import time
    for _ in range(50):
        status = client.get(f"/api/batches/{batch['batch_id']}").json()
        if status.get("status") == "awaiting_review":
            break
        time.sleep(0.05)
    assert status["status"] == "awaiting_review"
    assert client.post(f"/api/batches/{batch['batch_id']}/mask", json={"reviewed": False}).status_code == 409
    executed = client.post(f"/api/batches/{batch['batch_id']}/mask", json={"reviewed": True}).json()
    assert executed["status"] == "queued"

def test_file_mask_rejects_filename_extension_change():
    analyzed = client.post("/api/files/analyze", files={"file": ("check.txt", "电话13800138000".encode(), "text/plain")}).json()
    response = client.post("/api/files/mask", json={"analysis_id": analyzed["analysis_id"], "filename": "check.pdf", "text": analyzed["text"], "entities": analyzed["entities"], "reviewed": True})
    assert response.status_code == 422

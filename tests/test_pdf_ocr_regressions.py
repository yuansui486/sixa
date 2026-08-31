"""PDF OCR and custom-rule API regressions using generated local fixtures."""

from __future__ import annotations

import io
import json
import time
import zipfile

import fitz
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw

from app.main import app, ocr_service


class _FakeOcrEngine:
    def ocr(self, _path: str, cls: bool = False):
        return [[
            [[[30, 30], [430, 30], [430, 90], [30, 90]], ["13800138000", 0.99]],
        ]]


def _scan_pdf(page_count: int = 2) -> bytes:
    image = Image.new("RGB", (500, 120), "white")
    ImageDraw.Draw(image).text((30, 35), "13800138000", fill="black")
    image_bytes = io.BytesIO()
    image.save(image_bytes, format="PNG")
    document = fitz.open()
    for _ in range(page_count):
        page = document.new_page(width=500, height=120)
        page.insert_image(page.rect, stream=image_bytes.getvalue())
    data = document.tobytes()
    document.close()
    return data


def test_scanned_pdf_ocr_uses_page_offsets_and_redacts_boxes(monkeypatch):
    # Avoid downloading a model in the unit suite while exercising the exact
    # worker/extraction protocol used by the real Paddle service.
    monkeypatch.setattr(ocr_service, "engine", _FakeOcrEngine())
    monkeypatch.setattr(ocr_service, "_load_requested", True)
    with TestClient(app) as client:
        response = client.post(
            "/api/files/analyze",
            files={"file": ("scan.pdf", _scan_pdf(), "application/pdf")},
        )
        assert response.status_code == 200, response.text
        payload = response.json()
        assert payload["page_count"] == 2
        assert len(payload["boxes"]) == 2
        assert {box["page"] for box in payload["boxes"]} == {1, 2}
        assert len(payload["entities"]) == 2
        assert payload["entities"][1]["start"] > payload["entities"][0]["end"]
        assert all(entity["bbox"]["width"] > 0 for entity in payload["entities"])

        masked = client.post(
            "/api/files/mask",
            json={
                "analysis_id": payload["analysis_id"],
                "filename": "scan.pdf",
                "text": payload["text"],
                "entities": payload["entities"],
                "boxes": payload["boxes"],
                "reviewed": True,
            },
        )
        assert masked.status_code == 200, masked.text
        result = fitz.open(stream=client.get(masked.json()["artifact"]).content, filetype="pdf")
        assert all("13800138000" not in page.get_text() for page in result)
        assert all("138****8000" in page.get_text() for page in result)
        assert all("__MASKED_OCR_" not in page.get_text() for page in result)
        result.close()


def test_file_api_deselected_ocr_entity_does_not_mask_parent_region(monkeypatch):
    monkeypatch.setattr(ocr_service, "engine", _FakeOcrEngine())
    monkeypatch.setattr(ocr_service, "_load_requested", True)
    image = Image.new("RGB", (500, 120), "white")
    source = io.BytesIO()
    image.save(source, format="PNG")
    with TestClient(app) as client:
        analyzed = client.post(
            "/api/files/analyze",
            files={"file": ("selection.png", source.getvalue(), "image/png")},
        )
        assert analyzed.status_code == 200, analyzed.text
        payload = analyzed.json()
        assert len(payload["entities"]) == len(payload["boxes"]) == 1
        payload["entities"][0]["selected"] = False
        payload["boxes"][0]["selected"] = True
        masked = client.post(
            "/api/files/mask",
            json={
                "analysis_id": payload["analysis_id"],
                "filename": "selection.png",
                "entities": payload["entities"],
                "boxes": payload["boxes"],
                "reviewed": True,
            },
        )
        assert masked.status_code == 200, masked.text
        artifact = Image.open(io.BytesIO(client.get(masked.json()["artifact"]).content)).convert("RGB")
        assert artifact.getpixel((100, 60)) == (255, 255, 255)
        report = client.get(masked.json()["report"]).json()
        assert report["selected_entity_count"] == 0
        assert report["selected_box_count"] == 0


def test_batch_report_uses_effective_ocr_selection(monkeypatch):
    monkeypatch.setattr(ocr_service, "engine", _FakeOcrEngine())
    monkeypatch.setattr(ocr_service, "_load_requested", True)
    image = Image.new("RGB", (500, 120), "white")
    source = io.BytesIO()
    image.save(source, format="PNG")
    with TestClient(app) as client:
        created = client.post(
            "/api/batches",
            files={"files": ("batch-selection.png", source.getvalue(), "image/png")},
            data={"options": "{}"},
        )
        assert created.status_code == 200, created.text
        batch_id = created.json()["batch_id"]
        for _ in range(200):
            state = client.get(f"/api/batches/{batch_id}").json()
            if state.get("status") == "awaiting_review":
                break
            time.sleep(0.02)
        assert state["status"] == "awaiting_review"
        reviewed = state["items"]
        assert len(reviewed[0]["entities"]) == len(reviewed[0]["boxes"]) == 1
        reviewed[0]["entities"][0]["selected"] = False
        reviewed[0]["boxes"][0]["selected"] = True
        execute = client.post(
            f"/api/batches/{batch_id}/mask",
            json={"reviewed": True, "items": reviewed},
        )
        assert execute.status_code == 200, execute.text
        for _ in range(200):
            state = client.get(f"/api/batches/{batch_id}").json()
            if state.get("status") in {"completed", "partial", "failed"}:
                break
            time.sleep(0.02)
        assert state["status"] == "completed", state
        assert state["items"][0]["entities"][0]["selected"] is False
        assert state["items"][0]["boxes"][0]["selected"] is False
        archive_response = client.get(f"/api/batches/{batch_id}/download")
        assert archive_response.status_code == 200
        with zipfile.ZipFile(io.BytesIO(archive_response.content)) as archive:
            report = json.loads(archive.read("report.json"))
            assert report["items"][0]["boxes"][0]["selected"] is False
            image_name = next(name for name in archive.namelist() if name.endswith("_masked.png"))
            artifact = Image.open(io.BytesIO(archive.read(image_name))).convert("RGB")
            assert artifact.getpixel((100, 60)) == (255, 255, 255)


def test_custom_rule_replacement_is_applied_by_text_api():
    with TestClient(app) as client:
        created = client.post(
            "/api/rules",
            json={
                "name": "fixture project",
                "kind": "word",
                "pattern": "Project-LOCAL",
                "entity_type": "PROJECT_FIXTURE",
                "replacement": "[项目]",
            },
        )
        assert created.status_code == 200, created.text
        rule_id = created.json()["id"]
        try:
            analyzed = client.post("/api/text/analyze", json={"text": "Project-LOCAL"})
            assert analyzed.status_code == 200, analyzed.text
            entities = analyzed.json()["entities"]
            assert entities and entities[0]["replacement"] == "[项目]"
            masked = client.post(
                "/api/text/mask",
                json={"text": "Project-LOCAL", "entities": entities},
            )
            assert masked.status_code == 200, masked.text
            assert masked.json()["masked_text"] == "[项目]"
        finally:
            client.delete(f"/api/rules/{rule_id}")


def test_pdf_manual_box_is_masked_even_without_text_entity():
    document = fitz.open()
    page = document.new_page(width=200, height=120)
    page.insert_text((20, 60), "VISIBLE-SECRET")
    raw = document.tobytes()
    document.close()
    with TestClient(app) as client:
        analyzed = client.post(
            "/api/files/analyze",
            files={"file": ("manual.pdf", raw, "application/pdf")},
        )
        assert analyzed.status_code == 200, analyzed.text
        payload = analyzed.json()
        # The box covers the center of the generated page and is intentionally
        # independent of entity recognition.
        masked = client.post(
            "/api/files/mask",
            json={
                "analysis_id": payload["analysis_id"],
                "filename": "manual.pdf",
                "text": payload["text"],
                "entities": payload["entities"],
                "boxes": [{"page": 1, "x": 0.0, "y": 0.25, "width": 1.0, "height": 0.5, "selected": True}],
                "reviewed": True,
            },
        )
        assert masked.status_code == 200, masked.text
        output = client.get(masked.json()["artifact"]).content
        result = fitz.open(stream=output, filetype="pdf")
        assert "VISIBLE-SECRET" not in result[0].get_text()
        result.close()


def test_legacy_document_endpoint_supports_xlsx_and_legacy_image_keeps_format():
    from openpyxl import Workbook, load_workbook

    workbook = Workbook()
    workbook.active["A1"] = "电话 13800138000"
    source = io.BytesIO()
    workbook.save(source)
    with TestClient(app) as client:
        analyzed = client.post(
            "/api/document/analyze",
            files={"file": ("legacy.xlsx", source.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
        assert analyzed.status_code == 200, analyzed.text
        payload = analyzed.json()
        masked = client.post(
            "/api/document/mask",
            json={"analysis_id": payload["analysis_id"], "filename": "legacy.xlsx", "text": payload["text"], "entities": payload["entities"]},
        )
        assert masked.status_code == 200, masked.text
        result = load_workbook(io.BytesIO(client.get(masked.json()["artifact"]).content), data_only=False)
        assert "13800138000" not in str(result.active["A1"].value)

        image = Image.new("RGB", (80, 50), "white")
        image_bytes = io.BytesIO()
        image.save(image_bytes, format="JPEG")
        analyzed_image = client.post(
            "/api/image/analyze",
            files={"file": ("legacy.jpg", image_bytes.getvalue(), "image/jpeg")},
        )
        assert analyzed_image.status_code == 200, analyzed_image.text
        masked_image = client.post(
            "/api/image/mask",
            json={"analysis_id": analyzed_image.json()["analysis_id"], "boxes": [{"x": 0.1, "y": 0.1, "width": 0.5, "height": 0.5}]},
        )
        assert masked_image.status_code == 200, masked_image.text
        artifact = client.get(masked_image.json()["artifact"])
        assert artifact.status_code == 200
        assert Image.open(io.BytesIO(artifact.content)).format == "JPEG"

"""Disk-backed end-to-end fixtures for the local web workflow.

The adapters have focused unit tests, but a usable application also needs the
multipart upload, task persistence, artifact download, and post-download
parsers to work together.  These tests deliberately write source files to the
pytest temporary directory and upload those files through FastAPI's HTTP
surface rather than passing only synthetic in-memory payloads.
"""

from __future__ import annotations

import io
import time
import zipfile
from pathlib import Path

import fitz
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageChops, ImageDraw

from app.main import app


def _upload(client: TestClient, path: Path, content_type: str) -> dict:
    with path.open("rb") as source:
        response = client.post(
            "/api/files/analyze",
            files={"file": (path.name, source, content_type)},
        )
    assert response.status_code == 200, response.text
    return response.json()


def _mask(client: TestClient, analysis: dict, *, boxes: list[dict] | None = None) -> dict:
    response = client.post(
        "/api/files/mask",
        json={
            "analysis_id": analysis["analysis_id"],
            "filename": analysis["filename"],
            "text": analysis.get("text", ""),
            "entities": analysis.get("entities", []),
            "boxes": boxes if boxes is not None else analysis.get("boxes", []),
            "reviewed": True,
        },
    )
    assert response.status_code == 200, response.text
    artifact = client.get(response.json()["artifact"])
    assert artifact.status_code == 200
    return {**response.json(), "bytes": artifact.content}


def test_disk_backed_office_pdf_image_workflow(tmp_path: Path) -> None:
    docx_path = tmp_path / "客户资料.docx"
    document = Document()
    document.add_paragraph("联系电话 13800138000，邮箱 test@example.com")
    document.save(docx_path)

    xlsx_path = tmp_path / "客户资料.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "电话 13800138000"
    workbook.save(xlsx_path)

    pdf_path = tmp_path / "客户资料.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "电话 13800138000")
    pdf.save(pdf_path)
    pdf.close()

    image_path = tmp_path / "客户资料.png"
    original_image = Image.new("RGB", (240, 120), "white")
    # Use a non-default mask color so the assertion detects an actual pixel
    # change.  A black source region would be indistinguishable from the
    # default black solid policy even when the region was correctly selected.
    ImageDraw.Draw(original_image).rectangle((20, 20, 180, 95), fill=(35, 120, 210))
    original_image.save(image_path)

    client = TestClient(app)
    docx_analysis = _upload(
        client,
        docx_path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    docx_result = _mask(client, docx_analysis)
    masked_docx = Document(io.BytesIO(docx_result["bytes"]))
    assert "13800138000" not in "\n".join(p.text for p in masked_docx.paragraphs)

    xlsx_analysis = _upload(
        client,
        xlsx_path,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    xlsx_result = _mask(client, xlsx_analysis)
    masked_xlsx = load_workbook(io.BytesIO(xlsx_result["bytes"]), data_only=False)
    assert "13800138000" not in str(masked_xlsx.active["A1"].value)

    pdf_analysis = _upload(client, pdf_path, "application/pdf")
    pdf_result = _mask(client, pdf_analysis)
    masked_pdf = fitz.open(stream=pdf_result["bytes"], filetype="pdf")
    assert "13800138000" not in masked_pdf[0].get_text()
    masked_pdf.close()

    image_analysis = _upload(client, image_path, "image/png")
    image_result = _mask(
        client,
        image_analysis,
        boxes=[
            {
                "x": 20 / 240,
                "y": 20 / 120,
                "width": 160 / 240,
                "height": 75 / 120,
                "selected": True,
            }
        ],
    )
    masked_image = Image.open(io.BytesIO(image_result["bytes"])).convert("RGB")
    assert ImageChops.difference(original_image, masked_image).getbbox() is not None

    tasks = client.get("/api/tasks").json()["tasks"]
    ids = {item["id"] for item in tasks}
    assert {docx_analysis["analysis_id"], xlsx_analysis["analysis_id"], pdf_analysis["analysis_id"], image_analysis["analysis_id"]} <= ids


def test_disk_backed_reversible_and_batch_artifacts(tmp_path: Path) -> None:
    source = tmp_path / "reversible.txt"
    source.write_text("电话 13800138000", encoding="utf-8")
    client = TestClient(app)
    analysis = _upload(client, source, "text/plain")
    response = client.post(
        "/api/files/mask",
        json={
            "analysis_id": analysis["analysis_id"],
            "filename": source.name,
            "text": analysis["text"],
            "entities": analysis["entities"],
            "reviewed": True,
            "reversible": True,
            "password": "local-fixture-password",
        },
    )
    assert response.status_code == 200, response.text
    task_id = response.json()["task_id"]
    restored = client.post(
        f"/api/tasks/{task_id}/restore",
        json={"password": "local-fixture-password"},
    )
    assert restored.status_code == 200, restored.text
    restored_bytes = client.get(restored.json()["artifact"])
    assert restored_bytes.status_code == 200
    assert restored_bytes.content == source.read_bytes()

    first = tmp_path / "one.txt"
    second = tmp_path / "two.pdf"
    first.write_text("电话 13800138000", encoding="utf-8")
    pdf = fitz.open()
    pdf.new_page().insert_text((72, 72), "电话 13800138000")
    pdf.save(second)
    pdf.close()
    with first.open("rb") as first_stream, second.open("rb") as second_stream:
        created = client.post(
            "/api/batches",
            files=[
                ("files", (first.name, first_stream, "text/plain")),
                ("files", (second.name, second_stream, "application/pdf")),
            ],
            data={"options": "{}"},
        )
    assert created.status_code == 200, created.text
    batch_id = created.json()["batch_id"]
    for _ in range(200):
        status = client.get(f"/api/batches/{batch_id}").json()
        if status.get("status") == "awaiting_review":
            break
        time.sleep(0.03)
    assert status["status"] == "awaiting_review"
    executed = client.post(
        f"/api/batches/{batch_id}/mask",
        json={"reviewed": True, "items": status.get("items", [])},
    )
    assert executed.status_code == 200, executed.text
    for _ in range(300):
        status = client.get(f"/api/batches/{batch_id}").json()
        if status.get("status") in {"completed", "partial", "failed", "cancelled"}:
            break
        time.sleep(0.03)
    assert status["status"] == "completed", status
    archive = client.get(f"/api/batches/{batch_id}/download")
    assert archive.status_code == 200
    with zipfile.ZipFile(io.BytesIO(archive.content)) as bundle:
        names = set(bundle.namelist())
        assert "one_masked.txt" in names
        assert "two_masked.pdf" in names
        assert b"13800138000" not in bundle.read("one_masked.txt")
